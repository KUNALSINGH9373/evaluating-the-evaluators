#!/usr/bin/env python3
"""Surface candidate illustrative rows for each coding value, split by severity."""
import openpyxl, datetime, os, textwrap

ROOT = os.path.expanduser("~/MATS/Research/AISI_Evals")
ws = openpyxl.load_workbook(os.path.join(ROOT, "dataset", "AISIEVAL.xlsx"), data_only=True)["AISIEVAL"]
hdr = [c.value for c in ws[1]]
ix = {h: i + 1 for i, h in enumerate(hdr) if h}


def n(v):
    if v is None: return ""
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime("%Y-%m-%d")
    return str(v).strip()


R = [{h: n(ws.cell(r, ix[h]).value) for h in ix} for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
A = [r for r in R if r["Eval? (trackable)"] == "yes" and r["Action Trackable?"] == "yes"]
w = lambda s, k=150: textwrap.shorten(s.replace("\n", " "), k, placeholder=" …")

def show(col, val, sev, k=3):
    g = [r for r in A if r[col] == val and r["Severity (C1/C2) majority"] == sev]
    # prefer rows with a concise, self-contained response/policy text
    key = {"Action Level": "Company Response", "Policy Level": "Policy Response",
           "Attribution": "Company Response"}[col]
    g.sort(key=lambda r: (0 if 40 < len(r[key]) < 400 else 1, len(r[key])))
    print(f"\n{'='*104}\n{col} = {val}   ·   {sev}   ({len(g)} rows)")
    for r in g[:k]:
        print(f"  {r['Finding ID']:<28} {r['Institution'][:26]:<26} {r['Models / Systems'][:34]}")
        print(f"     finding : {w(r['Finding'], 145)}")
        if r[key]:
            print(f"     {key[:9]:<9}: {w(r[key], 170)}")
        if col == "Attribution" and r["Channel A Verbatim"]:
            print(f"     verbatim : {w(r['Channel A Verbatim'], 150)}")
        if col == "Policy Level" and r["Channel B Verbatim"]:
            print(f"     verbatim : {w(r['Channel B Verbatim'], 150)}")


for sev in ("C1", "C2"):
    for v in ("None", "Acknowledged", "Partial", "Substantive"):
        show("Action Level", v, sev)
print("\n\n" + "#" * 104 + "\nPOLICY LEVEL\n" + "#" * 104)
for sev in ("C1", "C2"):
    for v in ("No policy uptake identified", "Non-binding policy-related uptake", "Binding policy action"):
        show("Policy Level", v, sev)
print("\n\n" + "#" * 104 + "\nATTRIBUTION\n" + "#" * 104)
for sev in ("C1", "C2"):
    for v in ("Explicit attribution", "No explicit attribution", "No response located"):
        show("Attribution", v, sev, k=2)
