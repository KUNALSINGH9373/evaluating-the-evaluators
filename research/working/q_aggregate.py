#!/usr/bin/env python3
"""What are the Tier A rows coded with an 'aggregate' access type, and is the value right?"""
import openpyxl, os, re, datetime, collections

P = os.path.expanduser("~/Desktop/AISIEVAL.xlsx")
ws = openpyxl.load_workbook(P, data_only=True)[openpyxl.load_workbook(P).sheetnames[0]]
hdr = [c.value for c in ws[1]]
ix = {h: i + 1 for i, h in enumerate(hdr) if h}


def norm(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


R = [{h: norm(ws.cell(r, ix[h]).value) for h in ix} for r in range(2, ws.max_row + 1)
     if ws.cell(r, 1).value]
tier = lambda r: ("A" if r["Action Trackable?"] == "yes"
                  else "B" if r["Eval? (trackable)"] == "yes" else "C")

acol = [h for h in ix if re.search(r'access', str(h), re.I)]
print("access-ish columns:", acol)
for col in acol:
    print(f"\n=== {col} — full vocabulary ===")
    for k, v in collections.Counter(r[col] for r in R).most_common():
        a = sum(1 for r in R if r[col] == k and tier(r) == "A")
        print(f"  {v:>5} total · {a:>4} Tier A   {k!r}")

col = acol[0]
agg = [r for r in R if re.search(r'aggregat', r[col], re.I)]
print(f"\n=== rows whose access type mentions 'aggregate': {len(agg)} "
      f"({sum(1 for r in agg if tier(r)=='A')} Tier A) ===")
for r in sorted(agg, key=lambda r: r["Publication Date"]):
    print(f"\n--- {r['Finding ID']}  [Tier {tier(r)}]  {r['Publication Date']}  {r['Institution']}")
    print(f"    access  : {r[col]!r}")
    print(f"    models  : {r['Models / Systems']}")
    print(f"    report  : {r['Report Title']}")
    print(f"    finding : {r['Finding'][:420]}")
    print(f"    url     : {r['Source URL'][:110]}")
