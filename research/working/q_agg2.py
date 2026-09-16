#!/usr/bin/env python3
"""Was UK AISI's access to GPT-5.5 / Claude Mythos Preview pre- or post-deployment on 2026-05-13?

Also: how are the OTHER rows from the same report coded, and how are UK AISI's other
findings on the same two models coded? The right value should match its siblings.
"""
import openpyxl, os, re, datetime, collections

P = os.path.expanduser("~/Desktop/AISIEVAL.xlsx")
ws = openpyxl.load_workbook(P, data_only=True)[openpyxl.load_workbook(P).sheetnames[0]]
hdr = [c.value for c in ws[1]]
ix = {h: i + 1 for i, h in enumerate(hdr) if h}


def norm(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


R = [{h: norm(ws.cell(r, ix[h]).value) for h in ix} for r in range(2, ws.max_row + 1)
     if ws.cell(r, 1).value]
tier = lambda r: ("A" if r["Action Trackable?"] == "yes"
                  else "B" if r["Eval? (trackable)"] == "yes" else "C")

print("=== every row from the same report (Report ID UKAISI-2026-05) ===")
sib = [r for r in R if r["Report ID"].startswith("UKAISI-2026-05")]
for r in sorted(sib, key=lambda r: r["Finding ID"]):
    print(f"  {r['Finding ID']:<24} T{tier(r)}  access={r['Access Type']:<16} models={r['Models / Systems'][:44]}")

print("\n=== how is GPT-5.5 coded elsewhere in the corpus? ===")
for r in sorted([r for r in R if re.search(r'GPT-5\.5', r["Models / Systems"], re.I)],
                key=lambda r: r["Publication Date"]):
    print(f"  {r['Publication Date']}  {r['Finding ID']:<26} T{tier(r)}  access={r['Access Type']:<16} {r['Institution'][:24]}")

print("\n=== how is Claude Mythos Preview coded elsewhere? ===")
for r in sorted([r for r in R if re.search(r'Mythos Preview', r["Models / Systems"], re.I)],
                key=lambda r: r["Publication Date"]):
    print(f"  {r['Publication Date']}  {r['Finding ID']:<26} T{tier(r)}  access={r['Access Type']:<16} {r['Institution'][:24]}")

print("\n=== Access Type x Tier, to show how thin the Aggregate bucket is in Tier A ===")
ct = collections.Counter((r["Access Type"], tier(r)) for r in R)
for at in ("Pre-deployment", "Post-deployment", "Mixed", "Aggregate", "N/A"):
    print(f"  {at:<18} A={ct[(at,'A')]:>4}  B={ct[(at,'B')]:>4}  C={ct[(at,'C')]:>4}")
