#!/usr/bin/env python3
"""13_institution_type_tree.png — LEFT-TO-RIGHT institution tree.

Owns this figure outright. It previously lived inside hero.py, where a stale copy of that file
twice reverted the tree to a top-down layout whose leaf labels collided into unreadable runs.
Left-to-right is the whole point: institution names are long, so they must read horizontally on
their own row instead of being crammed under a narrow node.
"""
import os, collections, datetime, textwrap
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
import openpyxl

OUT = os.path.expanduser("~/MATS/Research/AISI_Evals/charts")
plt.rcParams.update({"figure.dpi": 200, "savefig.dpi": 200, "font.family": "DejaVu Sans",
                     "figure.facecolor": "white", "axes.facecolor": "white",
                     "savefig.facecolor": "white"})
L1 = "#C6DDF2"
PRIM = ["Government", "Non-Profit (AIEF)", "Non-Profit (Independent)", "For-Profit"]
COL = {"Government": "#1F6FB2", "Non-Profit (AIEF)": "#1E9E63",
       "Non-Profit (Independent)": "#5B9BD5", "For-Profit": "#E8A80C"}
WIRE = "#9AA7B2"


def norm(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


ws = openpyxl.load_workbook(os.path.expanduser("~/MATS/Research/AISI_Evals/dataset/AISIEVAL.xlsx"), data_only=True)["AISIEVAL"]
hdr = [norm(c.value) for c in ws[1]]
R = [{h: norm(ws.cell(r, i).value) for i, h in enumerate(hdr, 1) if h}
     for r in range(2, ws.max_row + 1)]
R = [r for r in R if r.get("Finding ID")]


def prim(r):
    p = [x.strip() for x in r.get("Institution Type", "").split(";") if x.strip()]
    for k in PRIM:
        if k in p:
            return k
    return p[0] if p else "Other"


grp = collections.defaultdict(list)
for r in R:
    grp[prim(r)].append(r)
order = sorted(PRIM, key=lambda k: -len(grp[k]))

# ---- build the leaf list per type, then lay out top to bottom -------------------------
TOPN = 5
groups = []
for k in order:
    top = collections.Counter(r["Institution"] for r in grp[k]).most_common(TOPN)
    other = len(grp[k]) - sum(v for _, v in top)
    leaves = top + ([("Other institutions", other)] if other > 0 else [])
    groups.append((k, len(grp[k]), leaves))

ROW = 1.0            # one unit of height per leaf row
GAP = 0.9            # blank rows between type groups
total = sum(len(g[2]) for g in groups) * ROW + GAP * (len(groups) - 1)

fig_h = 0.62 * total + 2.4
fig, ax = plt.subplots(figsize=(17, fig_h))
ax.set_xlim(0, 100)
ax.set_ylim(-1.2, total + 2.6)
ax.axis("off")
ax.invert_yaxis()

# geometry, all left-to-right
RX, RW = 1.0, 13.0          # root
TX, TW = 21.0, 21.0         # type nodes
CX, CW = 50.0, 7.2          # leaf count chips
NX = 58.6                   # leaf names start here

ax.text(RX, -0.75, "Who reports findings, by institution type",
        fontsize=30, color="#111111", va="bottom", ha="left")
ax.text(RX, 0.05, "Every finding traces to exactly one institution type. Node size is schematic, not proportional.",
        fontsize=15, color="#666666", va="bottom", ha="left")

y = 1.5
spans = []
for k, cnt, leaves in groups:
    h = len(leaves) * ROW
    spans.append((k, cnt, leaves, y, y + h))
    y += h + GAP

root_mid = (spans[0][3] + spans[-1][4]) / 2

# root
ax.add_patch(FancyBboxPatch((RX, root_mid - 1.15), RW, 2.3,
                            boxstyle="round,pad=0,rounding_size=0.28",
                            facecolor=L1, edgecolor="none"))
ax.text(RX + RW / 2, root_mid - 0.34, f"{len(R):,}", ha="center", va="center",
        fontsize=27, color="#111111")
ax.text(RX + RW / 2, root_mid + 0.58, "all findings", ha="center", va="center",
        fontsize=15, color="#111111")

# spine from root to each type node
mids = [(a + b) / 2 for _, _, _, a, b in spans]
ax.plot([RX + RW + 1.6, RX + RW + 1.6], [min(mids), max(mids)], color=WIRE, lw=1.6)
ax.plot([RX + RW, RX + RW + 1.6], [root_mid, root_mid], color=WIRE, lw=1.6)

for (k, cnt, leaves, y0, y1), mid in zip(spans, mids):
    ax.plot([RX + RW + 1.6, TX], [mid, mid], color=WIRE, lw=1.6)
    ax.add_patch(FancyBboxPatch((TX, mid - 1.15), TW, 2.3,
                                boxstyle="round,pad=0,rounding_size=0.28",
                                facecolor=COL[k], edgecolor="none"))
    ax.text(TX + 1.1, mid - 0.34, f"{cnt}", ha="left", va="center",
            fontsize=23, color="white")
    ax.text(TX + 1.1, mid + 0.62, k, ha="left", va="center", fontsize=14.5, color="white")

    # spine from type node out to its leaves
    lys = [y0 + ROW * (i + 0.5) for i in range(len(leaves))]
    bx = TX + TW + 2.2
    ax.plot([bx, bx], [min(lys), max(lys)], color=WIRE, lw=1.3)
    ax.plot([TX + TW, bx], [mid, mid], color=WIRE, lw=1.3)
    for (nm, v), ly in zip(leaves, lys):
        ax.plot([bx, CX], [ly, ly], color=WIRE, lw=1.3)
        ax.add_patch(FancyBboxPatch((CX, ly - 0.36), CW, 0.72,
                                    boxstyle="round,pad=0,rounding_size=0.16",
                                    facecolor=L1, edgecolor="none"))
        ax.text(CX + CW / 2, ly, f"{v}", ha="center", va="center",
                fontsize=15.5, color="#111111")
        ax.text(NX, ly, nm, ha="left", va="center", fontsize=15, color="#222222")

ax.text(RX, total + 2.35,
        "Institution Type field. Compound values (e.g. \"Government;Lab\") fold into their "
        "primary type.",
        fontsize=13.5, color="#333333", va="top", ha="left")

p = os.path.join(OUT, "13_institution_type_tree.png")
fig.savefig(p, bbox_inches="tight", pad_inches=0.32)
plt.close(fig)
print(f"wrote {p}  (left-to-right, {sum(len(g[2]) for g in groups)} leaf rows)")
for k, cnt, leaves in groups:
    print(f"  {k:<26} {cnt:>4}   " + " · ".join(f"{nm} {v}" for nm, v in leaves))

# The old bar version of this chart is superseded by the tree; keep it from reappearing.
old = os.path.join(OUT, "13_institution_type.png")
if os.path.exists(old):
    os.remove(old)
    print("  removed superseded 13_institution_type.png")
