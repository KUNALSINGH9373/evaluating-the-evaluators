#!/usr/bin/env python3
"""Make the deck open in Keynote.

DEFECT: python-pptx creates ppt/notesMasters/notesMaster1.xml and the presentation->notesMaster
relationship (rId8), but never adds <p:notesMasterIdLst> to ppt/presentation.xml. The notes
master is therefore referenced by every notesSlide while being undeclared by the presentation
itself. PowerPoint tolerates it; Keynote validates against ECMA-376 and refuses to open the file.

FIX: insert <p:notesMasterIdLst><p:notesMasterId r:id="rId8"/></p:notesMasterIdLst> in its
schema-mandated position -- CT_Presentation orders children
sldMasterIdLst, notesMasterIdLst, handoutMasterIdLst, sldIdLst, sldSz, notesSz, ... -- so it goes
immediately after </p:sldMasterIdLst>. Every other part is copied through byte-for-byte.
"""
import os, re, shutil, zipfile, time
from lxml import etree

P = os.path.expanduser("~/Desktop/AISIEVAL_10min.pptx")
TMP = P + ".repair"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NM = NS_R + "/notesMaster"

zin = zipfile.ZipFile(P)
rels = zin.read("ppt/_rels/presentation.xml.rels").decode()
m = re.search(r'Id="(rId\d+)"[^>]*Type="[^"]*/notesMaster"', rels)
if not m:
    raise SystemExit("no presentation->notesMaster relationship; a different fix is needed")
rid = m.group(1)
print(f"notesMaster relationship: {rid}")

pres = zin.read("ppt/presentation.xml").decode()
if "notesMasterIdLst" in pres:
    print("notesMasterIdLst already present — nothing to do")
    raise SystemExit(0)

inject = (f'<p:notesMasterIdLst><p:notesMasterId r:id="{rid}"/></p:notesMasterIdLst>')
new, n = re.subn(r'(</p:sldMasterIdLst>)', r'\1' + inject, pres, count=1)
if n != 1:
    raise SystemExit("could not locate </p:sldMasterIdLst>")
etree.fromstring(new.encode())                                  # must stay well-formed
order = [x for x in re.findall(r'<p:(\w+)', new) if x != "presentation"]
seq = [x for x in order if x.endswith("IdLst") or x in ("sldSz", "notesSz", "defaultTextStyle")]
print("child order now:", " -> ".join(dict.fromkeys(seq)))
EXPECT = ["sldMasterIdLst", "notesMasterIdLst", "sldIdLst", "sldSz", "notesSz"]
got = [x for x in dict.fromkeys(seq) if x in EXPECT]
assert got == [x for x in EXPECT if x in got], f"schema order wrong: {got}"
print("schema order: OK")

with zipfile.ZipFile(TMP, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "ppt/presentation.xml":
            data = new.encode()
        zi = zipfile.ZipInfo(item.filename, date_time=item.date_time)
        zi.compress_type = item.compress_type
        zi.external_attr = item.external_attr
        zout.writestr(zi, data)
zin.close()
shutil.copy2(P, P.replace(".pptx", f".backup-{time.strftime('%Y%m%d-%H%M%S')}.pptx"))
os.replace(TMP, P)
print(f"\nrepaired {P}  ({os.path.getsize(P):,} bytes)")

# --- reopen and re-verify -------------------------------------------------------------
from pptx import Presentation
z = zipfile.ZipFile(P)
assert z.testzip() is None
p2 = z.read("ppt/presentation.xml").decode()
print("notesMasterIdLst declared:", "notesMasterIdLst" in p2)
prs = Presentation(P)
print(f"reopens cleanly: {len(prs.slides._sldIdLst)} slides")
print("notes intact on all slides:",
      all(s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip() for s in prs.slides))
bad = [n for n in z.namelist() if n.endswith((".xml", ".rels"))]
for n in bad:
    etree.fromstring(z.read(n))
print(f"all {len(bad)} XML parts well-formed")
