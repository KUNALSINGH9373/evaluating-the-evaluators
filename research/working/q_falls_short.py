#!/usr/bin/env python3
"""What exactly is the 118/147 figure made of, and is 'the standard' severity-relative?"""
import openpyxl, os, collections, datetime

P = os.path.expanduser("~/Desktop/AISIEVAL.xlsx")
ws = openpyxl.load_workbook(P, data_only=True)[openpyxl.load_workbook(P).sheetnames[0]]
hdr = [c.value for c in ws[1]]
ix = {h: i + 1 for i, h in enumerate(hdr) if h}


def norm(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


R = [{h: norm(ws.cell(r, ix[h]).value) for h in ix} for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
tier = lambda r: ("A" if r["Action Trackable?"] == "yes" else "B" if r["Eval? (trackable)"] == "yes" else "C")
A = [r for r in R if tier(r) == "A"]
H = [r for r in A if r["Severity (C1/C2) majority"] == "C1"]
n = len(H)

print(f"C1 headline population: {n}\n")
print("composition of the 'falls short' figure:")
short, ok = [], []
for r in H:
    (ok if r["Proportionality"] == "Proportionate" else short).append(r)
by = collections.Counter(r["Action Level"] for r in short)
for lvl in ("None", "Acknowledged", "Partial"):
    if by[lvl]:
        print(f"  {by[lvl]:>4}  Action Level = {lvl:<13} -> "
              f"{'Accountability gap (no action)' if lvl=='None' else 'Under-response (gap)'}")
print(f"  {'':>4}  {'-'*54}")
print(f"  {len(short):>4}  falls short of the standard        = {len(short)/n:.1%}")
print(f"  {len(ok):>4}  meets it (Proportionate)           = {len(ok)/n:.1%}")
print(f"  {n:>4}  total\n")

print("why 'the standard' is severity-relative — the same Action Level scores differently:")
print("   severity  Action Level    outcome")
MATRIX = {("C1", "Substantive"): "Proportionate", ("C1", "Partial"): "Under-response (gap)",
          ("C1", "Acknowledged"): "Under-response (gap)", ("C1", "None"): "Accountability gap (no action)",
          ("C2", "Substantive"): "Proportionate", ("C2", "Partial"): "Proportionate",
          ("C2", "Acknowledged"): "Under-response (gap)", ("C2", "None"): "Accountability gap (no action)"}
for sev in ("C1", "C2"):
    for lvl in ("Substantive", "Partial", "Acknowledged", "None"):
        out = MATRIX[(sev, lvl)]
        mark = "  <-- Partial is enough for C2, not for C1" if (lvl == "Partial") else ""
        print(f"   {sev:<9} {lvl:<15} {out}{mark}")

C2 = [r for r in A if r["Severity (C1/C2) majority"] == "C2"]
c2short = sum(1 for r in C2 if r["Proportionality"] != "Proportionate")
print(f"\nfor comparison, the same measure over the {len(C2)} C2 rows: "
      f"{c2short}/{len(C2)} = {c2short/len(C2):.1%} fall short")
print(f"over all {len(A)} Tier A rows: "
      f"{sum(1 for r in A if r['Proportionality'] != 'Proportionate')}/{len(A)} = "
      f"{sum(1 for r in A if r['Proportionality'] != 'Proportionate')/len(A):.1%}")
