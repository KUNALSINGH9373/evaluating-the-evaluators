import openpyxl, re, json, collections
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
P='/Users/kunalsingh/Desktop/v11_FINAL.xlsx'
wb=openpyxl.load_workbook(P); ws=wb['v11 findings']
hdr=[c.value for c in ws[1]]
if 'What to check' not in hdr:
    ws.insert_cols(3); ws.cell(1,3,'What to check')
hdr=[c.value for c in ws[1]]; ix={h:i+1 for i,h in enumerate(hdr) if h}
G=lambda r,k: ws.cell(r,ix[k]).value or ''
S=lambda r,k,v: ws.cell(r,ix[k]).__setattr__('value',v)
tierA={x['row']:x for x in json.load(open('tierA.json'))}   # rows unchanged: insert_cols doesn't shift rows

PH=re.compile(r'-XX-|-\d{4}-00|^CIP-WEVAL')
for r in range(2,ws.max_row+1):
    st=G(r,'Status'); chk=[]
    if G(r,'Action Trackable?')=='yes':
        fl=tierA.get(r,{}).get('flags',[])
        chk += fl if fl else ['Spot-check only: confirm the quote is verbatim and the split/club call is right.']
    if st=='REVIEW-DATE':
        d=str(G(r,'Publication Date')).strip(); note=str(G(r,'Review note'))
        if not d: chk.append('DATE: no date on the source page. Blocked on the Weval date policy — one decision covers all 22.')
        elif re.fullmatch(r'\d{4}',d): chk.append('DATE: year-only (%s). Find month and day in the source.'%d)
        elif 'inferred' in note: chk.append('DATE: %s was inferred, not read. Read the real date off the source PDF.'%d)
        if PH.search(str(G(r,'Finding ID'))): chk.append('ID: placeholder Finding/Report ID — regenerate to PREFIX-YYYY-MM after the date lands.')
    S(r,'What to check',' | '.join(chk) if chk else '')

# --- Worklist sheet, grouped by report ---
if 'Worklist' in wb.sheetnames: del wb['Worklist']
wl=wb.create_sheet('Worklist',1)
wl.append(['Report / group','Evaluator','Rows','Sheet rows','What to check','Source URL'])
groups={}
for r in range(2,ws.max_row+1):
    if not G(r,'What to check'): continue
    key=(G(r,'Report ID'),G(r,'Institution'),str(G(r,'Report Title ')).strip(),G(r,'Source URL'))
    groups.setdefault(key,[]).append(r)
def prio(item):
    k,rs=item
    txt=' '.join(G(r,'What to check') for r in rs)
    p=0
    if 'SPLIT' in txt or 'TIER CONFLICT' in txt: p=0
    elif 'QUOTE:' in txt: p=1
    elif 'OVER-TIER' in txt: p=2
    elif 'NUMBERS' in txt: p=3
    elif 'DATE' in txt: p=4
    else: p=5
    return (p,-len(rs),k[1])
for (rid,inst,title,url),rs in sorted(groups.items(),key=prio):
    checks=sorted({c for r in rs for c in G(r,'What to check').split(' | ')})
    wl.append([rid,inst,len(rs),', '.join(str(x) for x in rs),'\n'.join('• '+c for c in checks),url])
hf=PatternFill('solid',start_color='D9D9D9',end_color='D9D9D9')
for c in range(1,7): wl.cell(1,c).fill=hf; wl.cell(1,c).font=Font(bold=True)
COL={'SPLIT':'F4CCCC','TIER CONFLICT':'F4CCCC','QUOTE:':'FCE4D6','OVER-TIER':'FFF2CC','NUMBERS':'FFF9E6','DATE':'DDEBF7'}
for r in range(2,wl.max_row+1):
    t=str(wl.cell(r,5).value)
    col=next((v for k,v in COL.items() if k in t),'E2EFDA')
    for c in range(1,7): wl.cell(r,c).fill=PatternFill('solid',start_color=col,end_color=col)
    wl.cell(r,5).alignment=Alignment(wrap_text=True,vertical='top')
    wl.cell(r,3).alignment=Alignment(horizontal='center')
for col,w in (('A',26),('B',30),('C',7),('D',18),('E',100),('F',44)): wl.column_dimensions[col].width=w
wl.freeze_panes='A2'; wl.auto_filter.ref=f'A1:F{wl.max_row}'
ws.freeze_panes='E2'
ws.column_dimensions['C'].width=90
for r in range(2,ws.max_row+1): ws.cell(r,3).alignment=Alignment(wrap_text=True,vertical='top')
ws.auto_filter.ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
wb.save(P)
print('worklist groups',wl.max_row-1,'| rows with a check',sum(1 for r in range(2,ws.max_row+1) if G(r,'What to check')))
