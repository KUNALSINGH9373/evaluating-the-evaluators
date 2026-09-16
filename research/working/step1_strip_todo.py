#!/usr/bin/env python3
"""Strip dead TODO scaffolding from Notes, and inventory the two remaining audit items.

The TODO strings are leftover generation scaffolding, not outstanding work: every row
carrying one already has its severity majority, and every Tier A row carrying one already
has Action Level and Policy Level. Verified before removal, and re-verified after.
"""
import openpyxl, os, shutil, time, re, datetime, collections, json

P = os.path.expanduser("~/Desktop/AISIEVAL.xlsx")
OUT = os.path.dirname(os.path.abspath(__file__))

PATS = [r'\s*TODO severity ensemble \+ Channel A/B/C\.?',
        r'\s*TODO — severity ensemble not yet run \(needs API keys, rulebook §[^)]*\)\.?']


def norm(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


shutil.copy2(P, P.replace(".xlsx", f".backup-{time.strftime('%Y%m%d-%H%M%S')}.xlsx"))
wb = openpyxl.load_workbook(P)
ws = wb[wb.sheetnames[0]]
hdr = [c.value for c in ws[1]]
ix = {h: i + 1 for i, h in enumerate(hdr) if h}
tier = lambda r: ("A" if r["Action Trackable?"] == "yes"
                  else "B" if r["Eval? (trackable)"] == "yes" else "C")

# --- safety gate: refuse to strip if any TODO row is genuinely incomplete -------------
rows = [{h: norm(ws.cell(r, ix[h]).value) for h in ix} | {"_r": r}
        for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
todo = [r for r in rows if "TODO" in r["Notes"].upper()]
incomplete = [r["Finding ID"] for r in todo
              if not r["Severity (C1/C2) majority"]
              or (tier(r) == "A" and not (r["Action Level"] and r["Policy Level"]))]
if incomplete:
    raise SystemExit(f"REFUSING: {len(incomplete)} TODO rows are genuinely incomplete: {incomplete[:10]}")
print(f"safety gate passed — all {len(todo)} TODO rows are complete; the text is scaffolding")

n = chars = 0
for r in rows:
    v = r["Notes"]
    if "TODO" not in v.upper():
        continue
    o = v
    for p in PATS:
        o = re.sub(p, '', o)
    o = re.sub(r'[ \t]+\n', '\n', o)
    o = re.sub(r'\n{3,}', '\n\n', o).strip()
    if o != v:
        ws.cell(r["_r"], ix["Notes"]).value = o
        n += 1
        chars += len(v) - len(o)
wb.save(P)
print(f"stripped {n} rows · {chars:,} chars removed")

# --- verify + inventory the two remaining items ---------------------------------------
ws2 = openpyxl.load_workbook(P, data_only=True)[wb.sheetnames[0]]
rows = [{h: norm(ws2.cell(r, ix[h]).value) for h in ix} for r in range(2, ws2.max_row + 1)
        if ws2.cell(r, 1).value]
left = [r["Finding ID"] for r in rows if "TODO" in r["Notes"].upper()]
print(f"rows still containing TODO: {len(left)} {left[:6]}")

A = [r for r in rows if tier(r) == "A"]
H = [r for r in A if r["Severity (C1/C2) majority"] == "C1"]
dated = lambda t: bool(re.search(r'20\d\d-\d\d-\d\d', t or ""))

nobat = [r for r in A if r["Action Level"] == "None" and not dated(r["Sources Checked (channel A)"])]
print(f"\n=== ITEM 2: {len(nobat)} rows at None with no dated battery ===")
for r in nobat:
    print(f"\n--- {r['Finding ID']}  ({r['Institution']}, pub {r['Publication Date']}, sev {r['Severity (C1/C2) majority']})")
    print(f"    models  : {r['Models / Systems']}")
    print(f"    report  : {r['Report Title']}")
    print(f"    url     : {r['Source URL']}")
    print(f"    finding : {r['Finding'][:340]}")
    print(f"    checked : {r['Sources Checked (channel A)'][:200]!r}")

sweep = sorted([r for r in H if r["Publication Date"] and r["Publication Date"][:4] < "2026"
                and r["Action Level"] == "None"], key=lambda r: r["Publication Date"])
print(f"\n=== ITEM 3: {len(sweep)} pre-2026 C1 rows coded None — the forward-in-time re-sweep ===")
by = collections.Counter()
for r in sweep:
    dev = r["Models / Systems"]
    key = ("OpenAI" if re.search(r'gpt|o[134]\b|chatgpt|codex|operator|sora', dev, re.I) else
           "Anthropic" if re.search(r'claude|opus|sonnet|haiku|fable|mythos', dev, re.I) else
           "Google" if re.search(r'gemini|gemma|palm', dev, re.I) else
           "Meta" if re.search(r'llama|prompt.?guard|muse', dev, re.I) else
           "xAI" if re.search(r'grok', dev, re.I) else
           "DeepSeek" if re.search(r'deepseek', dev, re.I) else
           "Mistral" if re.search(r'mistral|magistral', dev, re.I) else "other")
    by[key] += 1
print("by developer:", dict(by.most_common()))
print("by year:", dict(collections.Counter(r["Publication Date"][:4] for r in sweep)))
json.dump([{k: r[k] for k in ("Finding ID", "Institution", "Publication Date", "Models / Systems",
                              "Report Title", "Source URL", "Finding", "Domain")} for r in sweep],
          open(os.path.join(OUT, "sweep53.json"), "w"), indent=1)
print(f"\nwrote sweep53.json ({len(sweep)} rows)")
for r in sweep:
    print(f"  {r['Publication Date']}  {r['Finding ID']:<30} {r['Institution'][:26]:<26} {r['Models / Systems'][:44]}")
