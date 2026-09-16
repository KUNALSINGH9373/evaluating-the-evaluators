#!/usr/bin/env python3
"""Re-adjudicate the 11 rows individually: does the action address the IDENTIFIED PROBLEM?

Prints each finding beside the response that was cleared, so the call is made on the text
rather than on a bloc label. No writes.
"""
import openpyxl, os, csv, datetime, textwrap

P = os.path.expanduser("~/Desktop/AISIEVAL.xlsx")
CSV = os.path.expanduser("~/Documents/AISIEVAL_sensitivity_general_response.csv")
prior = {d["Finding ID"]: d for d in csv.DictReader(open(CSV, encoding="utf-8"))}
ws = openpyxl.load_workbook(P, data_only=True)[openpyxl.load_workbook(P).sheetnames[0]]
hdr = [c.value for c in ws[1]]
ix = {h: i + 1 for i, h in enumerate(hdr) if h}


def norm(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


w = lambda s, n=104, ind="      ": "\n".join(textwrap.wrap(s, n, initial_indent=ind, subsequent_indent=ind))
order = sorted(prior, key=lambda f: int(prior[f]["Lag (days)"] or 0))
for r in range(2, ws.max_row + 1):
    fid = ws.cell(r, 1).value
    if fid not in prior:
        continue
    prior[fid]["_finding"] = norm(ws.cell(r, ix["Finding"]).value)
    prior[fid]["_models"] = norm(ws.cell(r, ix["Models / Systems"]).value)

for fid in order:
    d = prior[fid]
    print("=" * 112)
    print(f"{fid}   sev {d['Severity']} · lag {d['Lag (days)']}d · was {d['Action Level']} / {d['Proportionality']}")
    print(f"  models  : {d['_models'][:96]}")
    print("  FINDING — the identified problem:")
    print(w(d["_finding"][:620]))
    print("  ACTION that was coded as the response:")
    print(w(d["Company Response"][:620]))
    if d.get("Channel A Verbatim"):
        print("  company's own words:")
        print(w('"' + d["Channel A Verbatim"][:400] + '"'))
