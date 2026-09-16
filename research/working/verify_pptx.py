#!/usr/bin/env python3
"""Re-open the deck and confirm Keynote will get editable text, images and presenter notes."""
import os, zipfile
from pptx import Presentation
from pptx.util import Emu

P = os.path.expanduser("~/Desktop/AISIEVAL_10min.pptx")
z = zipfile.ZipFile(P)
bad = z.testzip()
print(f"zip integrity: {'OK' if bad is None else 'CORRUPT at ' + bad}")
print(f"parts: {len(z.namelist())} · media: {sum(1 for n in z.namelist() if n.startswith('ppt/media/'))} "
      f"· notesSlides: {sum(1 for n in z.namelist() if 'notesSlide' in n)}")

prs = Presentation(P)
SW, SH = prs.slide_width, prs.slide_height
print(f"slide size: {SW/914400:.3f} x {SH/914400:.3f} in  (16:9 = {SW/SH:.4f})\n")
tot_txt = tot_pic = 0
ok = True
for i, s in enumerate(prs.slides, 1):
    txt = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    pics = [sh for sh in s.shapes if sh.shape_type == 13]
    tbls = [sh for sh in s.shapes if sh.has_table]
    words = sum(len(sh.text_frame.text.split()) for sh in txt)
    note = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
    tot_txt += len(txt)
    tot_pic += len(pics)
    off = []
    for sh in s.shapes:
        if sh.left is None or sh.top is None or sh.width is None:
            continue
        if sh.left < 0 or sh.top < 0 or sh.left + sh.width > SW + Emu(9525) \
           or sh.top + sh.height > SH + Emu(9525):
            off.append(sh.shape_type)
    if off:
        ok = False
    print(f"  slide {i}: {len(txt)} text boxes ({words} words) · {len(pics)} charts · "
          f"{len(tbls)} table(s) · notes {len(note.split())} words"
          + (f"  !! {len(off)} shape(s) off-canvas" if off else ""))
print(f"\ntotals: {tot_txt} editable text boxes · {tot_pic} chart images")
print(f"off-canvas overflow: {'none' if ok else 'SEE ABOVE'}")
missing = [i for i, s in enumerate(prs.slides, 1)
           if not (s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())]
print(f"slides missing presenter notes: {missing or 'none'}")
