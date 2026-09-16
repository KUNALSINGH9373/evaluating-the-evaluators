# -*- coding: utf-8 -*-
import openpyxl,shutil,time,collections
P='/Users/kunalsingh/Desktop/v11_FINAL.xlsx'
shutil.copy2(P,'/Users/kunalsingh/Desktop/v11_FINAL.backup-%s.xlsx'%time.strftime('%Y%m%d-%H%M%S'))
SD='2026-08-15'
LOG=("Channel B battery run %s. UK: Hansard API (Commons + Lords), exact-phrase verified against "
     "ContributionTextFull; gov.uk and NCSC site-restricted search. US: Federal Register API "
     "(exact phrase), congress.gov / nist.gov / commerce.gov / whitehouse.gov site-restricted search "
     "(Congress.gov API requires a key not held). International: International AI Safety Report 2026 "
     "(3 Feb 2026, backed by 30+ governments) full-text checked for this evaluator and finding. "
     "No official government source references this finding. ")%SD
FALSE=("Recorded invalidations: Federal Register exact-phrase hits for 'METR' (Medicare rules), "
       "'Scale AI' ('Content at Scale AI', a company name) and 'RAND Europe' (air-quality and "
       "fuel-economy rules) are substring or unrelated-entity matches, not policy uptake. Hansard "
       "hits for 'METR' (marginal effective tax rate) and 'Scale AI' ('scale AI capabilities') are "
       "likewise not references to the evaluator.")
UPTAKE={'APOLLO-2023-11-ALI1','OPENAI-2024-09-SELF-ALI1'}
UP=dict(level='Non-binding policy uptake',
  resp=("Ben Lake MP cited Apollo Research by name in the House of Commons AI Safety debate "
        "(Westminster Hall, 10 December 2025), referencing its finding that an OpenAI model "
        "attempted to deceive users and to disable monitoring mechanisms, as evidence that advanced "
        "models deploy techniques to avoid human control."),
  verb=("Apollo Research found examples of one of OpenAI's models trying to deceive users to "
        "accomplish its goals and, perhaps most worryingly, to disable monitoring mechanisms and guardrails."),
  evid=("https://hansard.parliament.uk/Commons/2025-12-10/debates/9F01B4B9-12CB-42E2-84E2-A65F7D30BFAF/AISafety "
        "(Hansard, Commons Westminster Hall, 10 Dec 2025; Ben Lake, Ceredigion Preseli, PC). Retrieved via the "
        "Hansard Search API and verified against ContributionTextFull."))
wb=openpyxl.load_workbook(P); ws=wb[wb.sheetnames[0]]
hdr=[c.value for c in ws[1]]; ix={h:i+1 for i,h in enumerate(hdr) if h}
G=lambda r,k: ws.cell(r,ix[k]).value
S=lambda r,k,v: ws.cell(r,ix[k]).__setattr__('value',v)
n=collections.Counter()
for r in range(2,ws.max_row+1):
    if G(r,'Action Trackable?')!='yes': continue
    fid=G(r,'Finding ID')
    if fid in UPTAKE:
        S(r,'Policy Level',UP['level']); S(r,'Policy Response',UP['resp'])
        S(r,'Channel B Verbatim',UP['verb'])
        S(r,'Channel B Evidence',UP['evid']+" NOTE: the single cited sentence covers both Apollo/OpenAI findings in this dataset - 'deceive users' maps to APOLLO-2023-11-ALI1 and 'disable monitoring mechanisms' to the o1 scheming evaluation in OPENAI-2024-09-SELF-ALI1; both are coded from it.")
        n[UP['level']]+=1
    else:
        S(r,'Policy Level','No policy uptake identified')
        S(r,'Channel B Evidence',LOG+FALSE)
        n['No policy uptake identified']+=1
wb.save(P)
print('Channel B coded:',dict(n))
tot=collections.Counter(G(r,'Policy Level') for r in range(2,ws.max_row+1) if G(r,'Action Trackable?')=='yes')
print('Tier A Policy Level:',dict(tot))
print('blank:',sum(1 for r in range(2,ws.max_row+1) if G(r,'Action Trackable?')=='yes' and not G(r,'Policy Level')))
