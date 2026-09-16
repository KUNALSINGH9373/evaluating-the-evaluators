#!/usr/bin/env python3
"""Is 'FAR.AI' in those response cells the COMPANY naming the evaluator, or my own commentary?"""
import openpyxl, os, re, datetime

P = os.path.expanduser("~/Desktop/AISIEVAL.xlsx")
ws = openpyxl.load_workbook(P, data_only=True)[openpyxl.load_workbook(P).sheetnames[0]]
hdr = [c.value for c in ws[1]]
ix = {h: i + 1 for i, h in enumerate(hdr) if h}
for r in range(2, ws.max_row + 1):
    fid = ws.cell(r, 1).value
    if fid not in ("FARAI-2024-08-JAI1", "UKAISI-2026-04-CYB11"):
        continue
    cr = str(ws.cell(r, ix["Company Response"]).value or "")
    print("=" * 100)
    print(f"{fid}   attr={ws.cell(r, ix['Attribution']).value!r}  lag={ws.cell(r, ix['Lag (days)']).value}")
    print(f"FULL RESPONSE CELL ({len(cr)} chars):\n{cr}\n")
