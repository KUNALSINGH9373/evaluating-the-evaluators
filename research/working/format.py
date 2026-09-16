#!/usr/bin/env python3
"""Re-apply tier colour banding and mark review rows in red."""
import os, shutil, time, collections, datetime, openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

P=os.path.expanduser("~/Desktop/AISIEVAL.xlsx"); D="2026-08-16"
TIER_FILL={"A":"FFFCE4E4","B":"FFFFF6DC","C":"FFEAF2FC"}   # A red-tint · B amber-tint · C blue-tint
HEADER="FF404040"; REVIEW_FILL="FFFFE0E0"; RED="FFC00000"

# a row needs review if any cell carries one of these markers
MARKERS={
 "UNRESOLVED":"unresolved inconsistency",
 "INTERPRETIVE FLAG":"Substantive/Acknowledged judgement",
 "INADMISSIBLE RESPONSE EVIDENCE":"response documented only by the evaluator",
 "CIRCULARITY FLAG":"Channel B self-announcement",
 "NOTE FOR REVIEW":"Tier A assignment questionable",
 "NOT VERIFIED":"attribution unverified",
 "developer-mismatch":"developer mismatch",
 "NOTE: the finding records":"developer mismatch",
}
def norm(v):
    if v is None: return ""
    if isinstance(v,(datetime.datetime,datetime.date)): return v.strftime("%Y-%m-%d")
    return str(v).strip()

shutil.copy2(P, P.replace(".xlsx",".backup-"+time.strftime("%Y%m%d-%H%M%S")+".xlsx"))
wb=openpyxl.load_workbook(P); ws=wb["AISIEVAL"]
hdr=[norm(c.value) for c in ws[1]]; ix={h:i for i,h in enumerate(hdr,1) if h}
ncol=len(hdr); last=get_column_letter(ncol)
def g(r,n): return norm(ws.cell(r,ix[n]).value)

# --- header -------------------------------------------------------------
for i in range(1,ncol+1):
    c=ws.cell(1,i)
    c.fill=PatternFill("solid",start_color=HEADER,end_color=HEADER)
    c.font=Font(bold=True,color="FFFFFFFF",size=11)
    c.alignment=Alignment(vertical="center",horizontal="left",wrap_text=True)
ws.row_dimensions[1].height=34

WIDE={"Finding ID":30,"Report ID":28,"Institution":30,"Institution Type":18,"Report Title":42,
 "Publication Date":14,"Domain":22,"Models / Systems":30,"Access Type":16,"Source URL":44,
 "Finding":72,"Finding Quote":60,"Company Response":52,"Channel A Verbatim":48,"Channel A Evidence":44,
 "Sources Checked (channel A)":54,"Policy Response":44,"Channel B Verbatim":44,"Channel B Evidence":54,
 "Media Outlets":34,"Academic Citations":30,"Social Highlights":30,"Channel C Verbatim":34,
 "Notes":80,"Finding Type":26,"Tags":22,"Scope":22,"Source Dataset":14}
for i,h in enumerate(hdr,1):
    ws.column_dimensions[get_column_letter(i)].width=WIDE.get(h,16)

# --- body ----------------------------------------------------------------
tiers=collections.Counter(); review=collections.Counter(); rows_flagged=[]
thin=Side(style="thin",color="FFD0D0D0")
for r in range(2, ws.max_row+1):
    if not g(r,"Finding ID"): continue
    t=("A" if g(r,"Action Trackable?")=="yes" else "B" if g(r,"Eval? (trackable)")=="yes" else "C")
    tiers[t]+=1
    fill=PatternFill("solid",start_color=TIER_FILL[t],end_color=TIER_FILL[t])
    for i in range(1,ncol+1):
        cell=ws.cell(r,i)
        cell.fill=fill
        cell.alignment=Alignment(vertical="top",wrap_text=False)
        cell.border=Border(bottom=thin)
        cell.font=Font(size=10)
    # tier column emphasis
    for col in ("Eval? (trackable)","Action Trackable?"):
        ws.cell(r,ix[col]).font=Font(size=10,bold=True)
    ws.cell(r,ix["Finding ID"]).font=Font(size=10,bold=True)
    # review markers
    blob=" ".join(g(r,h) for h in hdr if h)
    hits=sorted({lab for k,lab in MARKERS.items() if k in blob})
    if hits:
        n=g(r,"Notes")
        tag=f"[REVIEW NEEDED: {'; '.join(hits)}] "
        if not n.startswith("[REVIEW NEEDED"):
            ws.cell(r,ix["Notes"]).value=tag+n
        nc=ws.cell(r,ix["Notes"])
        nc.font=Font(size=10,bold=True,color=RED)
        nc.fill=PatternFill("solid",start_color=REVIEW_FILL,end_color=REVIEW_FILL)
        fid=ws.cell(r,ix["Finding ID"]); fid.font=Font(size=10,bold=True,color=RED)
        for h in hits: review[h]+=1
        rows_flagged.append((g(r,"Finding ID"),hits))

ws.freeze_panes="B2"
for name in list(ws.tables): del ws.tables[name]
tbl=Table(displayName="AISIEVAL", ref=f"A1:{last}{ws.max_row}")
tbl.tableStyleInfo=TableStyleInfo(name="TableStyleLight1",showRowStripes=False,showColumnStripes=False)
ws.add_table(tbl)
wb.save(P)

print("tier banding applied: " + " · ".join(f"{k} {tiers[k]}" for k in "ABC"))
print(f"  A = {TIER_FILL['A']} (red tint) · B = {TIER_FILL['B']} (amber) · C = {TIER_FILL['C']} (blue)")
print(f"\nrows marked REVIEW NEEDED (red Notes): {len(rows_flagged)}")
for k,v in review.most_common(): print(f"   {v:>3}  {k}")
print("\n  " + "\n  ".join(f"{f:<28} {'; '.join(h)}" for f,h in rows_flagged))
