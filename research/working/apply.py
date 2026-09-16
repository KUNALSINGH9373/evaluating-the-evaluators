import openpyxl, collections
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

P='/Users/kunalsingh/Desktop/v11_FINAL.xlsx'
wb=openpyxl.load_workbook(P); ws=wb['v11 findings']
hdr=[c.value for c in ws[1]]; ix={h:i+1 for i,h in enumerate(hdr) if h}
S=lambda r,k,v: ws.cell(r,ix[k]).__setattr__('value',v)
G=lambda r,k: ws.cell(r,ix[k]).value
STAMP=" §1 evaluation-vs-benchmark test applied 2026-08-15."

# ---------- 1. Row 8: Pattern Labs / o3-o4-mini -> reinstate, corrected ----------
S(8,'Status','READY')
S(8,'Eval? (trackable)','yes'); S(8,'Action Trackable?','no')
S(8,'Finding',"Pattern Labs evaluated o3 and o4-mini across Evasion, Network Attack Simulation, and Vulnerability Discovery and Exploitation challenges and concluded that, despite o3 showing improved cyberoffensive capability over o3-mini, limitations in recognising and recovering from tactical setbacks left it unable to solve hard challenges, so o3 would provide only limited assistance to a moderately skilled cyberoffensive operator.")
S(8,'Finding Quote',"These and other limitations contributed to the model's inability to solve hard challenges, and led Pattern Labs to conclude that despite the performance improvements, o3 would provide only limited assistance to a moderately skilled cyberoffensive operator.")
S(8,'Finding Type','capability-finding;reassuring-null;company-published')
S(8,'Review note',"Reinstated 2026-08-15. Prior exclusion ('reports no outcome') was incorrect — the Pattern Labs conclusion appears later in §3.9.3 than the original extraction read. Pattern Labs ran the evaluation and asserts the conclusion, so it passes the §1 evaluation-vs-benchmark test. Retiered A->B: the conclusion is reassuring, no danger threshold crossed.")
S(8,'Notes',"Pattern Labs RAN the evaluation and asserts the conclusion -> qualifies under §1." + STAMP + " Re-extracted from o3-and-o4-mini-system-card.pdf §3.9.3. TODO severity ensemble + Channel A/B/C.")

# ---------- 2. Rows failing the test -> mark for removal ----------
DROP={
 38:"Removed 2026-08-15 (§1 evaluation-vs-benchmark test). Meta ran Anthropic's Agentic Misalignment benchmark itself — the card's own words are 'we also evaluate the model in agentic contexts'. Apollo Research is cited only for the separate evaluation-awareness finding (META-2026-07-MUSESPARK-ALI1); attributing this row to Apollo was an extraction error.",
 39:"Removed 2026-08-15 (§1 evaluation-vs-benchmark test). Same Meta-run Agentic Misalignment table as ALI2; Apollo attribution is an extraction error. Subject model is Gemini 3.1 Pro, reported by Meta as a comparator.",
 144:"Removed 2026-08-15 (§1 evaluation-vs-benchmark test). Sonnet 5 card: the IPI benchmark is 'built in partnership with Gray Swan, the UK AI Security Institute, the US Center for AI Standards and Innovation, and other model developers', then 'we organized a competition' — Anthropic ran the exercise and asserts the result. Benchmark co-construction is not contribution.",
 165:"Removed 2026-08-15 (§1 named-evaluator test). Mythos Preview card §2.3.7 attributes this block to 'additional external partners' — unnamed. METR and Epoch AI are named in the same section but not for this finding. The recorded attribution to UK AISI is unverified; UK AISI is not named anywhere in the card.",
 166:"Removed 2026-08-15 (§1 named-evaluator test). Same unnamed 'additional external partners' block as AUT1 (Mythos Preview card §2.3.7).",
 177:"Removed 2026-08-15 (§1 evaluation-vs-benchmark test). 'After reviewing 141,006 evaluation runs ... we identified three incidents' — Anthropic ran the review of its own evaluation infrastructure and asserts the finding. Irregular is the environment in which the incidents occurred, not the evaluator asserting them.",
}
for rn,note in DROP.items():
    S(rn,'Status','REMOVED'); S(rn,'Review note',note)

# ---------- 3. Corrections ----------
S(66,'Finding',"Gray Swan ran prompt-injection attacks against Meta's Muse Spark using its private Agent Red Teaming (ART) benchmark before launch, and Muse Spark showed greater sensitivity to injection than the other models Gray Swan evaluated.")
S(66,'Review note',"Corrected 2026-08-15. Gray Swan RAN this evaluation — Muse Spark report §3.2.2.4: 'GraySwan additionally ran prompt injection attacks using their private Agent Red Teaming (ART) benchmark (Zou et al., 2025) before the Muse Spark launch.' Meta's own AgentDojo pass@1 figure of 11.7% has been removed from the Finding — it is a different benchmark, run by Meta, and did not belong in a Gray Swan row. Verify the quote and split/club, then sign off.")
S(66,'Notes',"Gray Swan RAN the evaluation (report §3.2.2.4) -> contribution, not benchmark usage." + STAMP + " Extracted 2026-08-14, corrected 2026-08-15. TODO severity ensemble + Channel A/B/C.")

S(298,'Institution','Gray Swan AI')
S(298,'Scope','third-party-evaluator')
S(298,'Tags','indirect-prompt-injection;art-benchmark;gray-swan;anthropic')
S(298,'Finding',"Gray Swan evaluated Claude Sonnet 4.6 on the Agent Red Teaming benchmark; Sonnet 4.6 improved significantly over Sonnet 4.5 and performed comparably to Opus 4.6 on indirect prompt-injection resistance.")
S(298,'Review note',"Attribution corrected 2026-08-15. Gray Swan ran the evaluation; UK AISI only co-developed the ART benchmark, which under the §1 test is not contribution. UK AISI removed from Institution and Scope changed government-AISI -> third-party-evaluator.")
S(298,'Notes',"Gray Swan RAN it -> contribution. UK AISI co-developed the ART benchmark only -> benchmark authorship is not contribution, so not government-AISI scope." + STAMP + " Improvement, no threshold crossed -> Tier B. Extracted 2026-08-14, corrected 2026-08-15. TODO severity ensemble + Channel A/B/C.")

S(391,'Finding',"METR ran a GPT-4o-based LLM agent on a suite of 86 long-horizon, multi-step, end-to-end tasks across 31 task families in virtual environments and found the model unable to robustly take autonomous actions: it typically completed individual substeps such as creating SSH keys or logging into VMs but did not carry tasks through to completion.")
S(391,'Finding Quote',"METR ran a GPT-4o-based simple LLM agent on a suite of long-horizon multi-step end-to-end tasks in virtual environments. ... GPT-4o was unable to robustly take autonomous actions. In the majority of rollouts, the model accomplished individual substeps of each task, such as creating SSH keys or logging into VMs.")
S(391,'Review note',"Quote corrected 2026-08-15. The previous quote spliced OpenAI's own line ('Select machine learning engineering tasks from METR: 0/10 trials' — OpenAI running METR's tasks) with METR's conclusion. Replaced with METR's own assessment from §4.1, where METR ran the agent. The 0/10 figure has been moved to Notes as an OpenAI-run comparator.")
S(391,'Notes',"METR RAN this assessment (GPT-4o system card §4.1 'METR assessment') -> qualifies under §1." + STAMP + " [OpenAI separately reported 0/10 on selected ML engineering tasks from METR in §3 — that is OpenAI running METR's task set, not a METR finding, and is excluded per the benchmark-usage rule.] Reassuring null -> Tier B. Extracted 2026-08-14, corrected 2026-08-15. TODO severity ensemble + Channel A/B/C.")

# ---------- 4. Redwood: SLEIGHT-Bench included; SHADE-Arena attribution ----------
for rn in (73,431):
    S(rn,'Notes',(G(rn,'Notes') or '').split(' Extracted')[0].rstrip('. ') +
      ". Authorship is 4:1 Anthropic:Redwood (Elle Najt and Colin Toft, Anthropic Fellows Program; Tyler Tracy, Redwood Research; Fabien Roger and Joe Benton, Anthropic) and the work was done inside the Anthropic Fellows Program. Retained by explicit ruling 2026-08-15: a named independent co-author who ran the work satisfies §1. Contrast SHADE-Arena (ANTHROPIC-2025-06-SELF-ALI1..4), a stated three-way collaboration between Anthropic, Scale AI and Redwood Research — a stronger case on the same rule." +
      STAMP + " TODO severity ensemble + Channel A/B/C.")
S(73,'Review note',"Tier A. Retained by ruling 2026-08-15 despite 4:1 Anthropic:Redwood authorship — Tyler Tracy (Redwood Research) is a named co-author who ran the work, which satisfies §1. Verify the verbatim quote and the split/club decision, then sign off.")
S(431,'Review note',"Retained by ruling 2026-08-15 despite 4:1 Anthropic:Redwood authorship — see Notes. Tier C methodology. No further action.")

for rn in (427,428,429,430):
    S(rn,'Institution','Redwood Research + Scale AI')
    S(rn,'Notes',(G(rn,'Notes') or '').split(' Extracted')[0].rstrip('. ') +
      ". Attribution corrected 2026-08-15: the post states 'This work was a collaboration between researchers at Anthropic, Scale AI (Xiang Deng and Chen Bo Calvin Zhang) and Redwood Research (Tyler Tracy and Buck Shlegeris)'. Scale AI was a named co-author and was missing from Institution." + STAMP + " TODO severity ensemble + Channel A/B/C.")
    S(rn,'Review note',"Attribution corrected 2026-08-15 — Scale AI added as named co-author alongside Redwood Research. Retained under §1: both are named independent co-authors who ran the work. Tier B/C, no further action.")

wb.save(P)
print('stage 1 saved')
c=collections.Counter(ws.cell(r,ix['Status']).value for r in range(2,ws.max_row+1))
print(dict(c))
