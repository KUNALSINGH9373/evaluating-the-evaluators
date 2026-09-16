import openpyxl, re, collections, json
wb=openpyxl.load_workbook('/Users/kunalsingh/Desktop/v11_FINAL.xlsx'); ws=wb['v11 findings']
hdr=[c.value for c in ws[1]]; ix={h:i+1 for i,h in enumerate(hdr) if h}
G=lambda r,k: ws.cell(r,ix[k]).value or ''
COMPANIES=['OpenAI','Anthropic','Google','DeepMind','Meta','Mistral','xAI','Grok','DeepSeek','Qwen','Alibaba','Cohere','Microsoft','Amazon','Nvidia','Llama','Gemini','Claude','GPT','o1','o3','o4']
FAMILY={'GPT':'OpenAI','o1':'OpenAI','o3':'OpenAI','o4':'OpenAI','OpenAI':'OpenAI','Claude':'Anthropic','Anthropic':'Anthropic',
 'Gemini':'Google','DeepMind':'Google','Google':'Google','Llama':'Meta','Meta':'Meta','Grok':'xAI','xAI':'xAI',
 'DeepSeek':'DeepSeek','Qwen':'Alibaba','Alibaba':'Alibaba','Mistral':'Mistral','Cohere':'Cohere','Microsoft':'Microsoft'}
THRESH=re.compile(r'threshold|uplift|danger|critical|crossed|exceed|first model|no other|never|unable to|failed to|catastroph|high capability|significant risk|jailbreak|bypass|exfiltrat|sabotage|deceiv|deception|scheming|misalign|refus|attack|vulnerab|escape|unauthoris|unauthoriz',re.I)
RANK=re.compile(r'outperform|ranked|highest|lowest|best of the|worst|compared with other|relative to other|top of|leaderboard|scored higher|better than',re.I)
rows=[]
for r in range(2,ws.max_row+1):
    if G(r,'Action Trackable?')!='yes': continue
    ft=str(G(r,'Finding Type')); f=str(G(r,'Finding')); q=str(G(r,'Finding Quote')); ms=str(G(r,'Models / Systems'))
    flags=[]
    for bad,msg in [('reassuring-null','TIER CONFLICT: tagged reassuring-null but Tier A — §4 excludes reassuring nulls'),
                    ('anonymised-model','TIER CONFLICT: tagged anonymised-model but Tier A — no accountable party'),
                    ('non-frontier','TIER CONFLICT: tagged non-frontier but Tier A'),
                    ('capability-trend','TIER CONFLICT: Finding Type is capability-trend — §4 puts trends in Tier B')]:
        if bad in ft: flags.append(msg)
    fams={FAMILY[k] for k in FAMILY if re.search(r'\b'+re.escape(k),ms)}
    if len(fams)>1: flags.append('SPLIT: Models/Systems spans %s — §5 forbids two companies on one row'%('+'.join(sorted(fams))))
    if not THRESH.search(f) and not THRESH.search(q):
        flags.append('OVER-TIER RISK: no threshold/harm language — reads as a bare score, which §4 excludes from Tier A')
    if RANK.search(f) and not THRESH.search(f):
        flags.append('OVER-TIER RISK: comparative ranking — §4 excludes rankings from Tier A')
    if len(q.strip())<70: flags.append('QUOTE: only %d chars — too short to carry the finding'%len(q.strip()))
    nums=set(re.findall(r'\d+\.?\d*%?',f)); qn=set(re.findall(r'\d+\.?\d*%?',q))
    miss=[n for n in nums if n not in qn and len(n)>1]
    if miss: flags.append('NUMBERS: %s in Finding but not in Quote — verify against source'%', '.join(sorted(miss)[:4]))
    rows.append(dict(row=r,fid=G(r,'Finding ID'),rid=G(r,'Report ID'),inst=G(r,'Institution'),
                     title=str(G(r,'Report Title ')),url=G(r,'Source URL'),status=G(r,'Status'),flags=flags))
json.dump(rows,open('tierA.json','w'),indent=1)
print('Tier A rows',len(rows))
print('clean (quote+split spot-check only):',sum(1 for x in rows if not x['flags']))
c=collections.Counter(f.split(':')[0] for x in rows for f in x['flags'])
for k,v in c.most_common(): print(' ',v,k)
