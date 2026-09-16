#!/usr/bin/env python3
"""Item 2: the three rows coded None with no dated Channel A battery.

Investigation found two of the three were under-coded, not merely undated, and the third
exposed a cross-row inconsistency that is the user's call to settle.
"""
import openpyxl, os, shutil, time, collections

P = os.path.expanduser("~/Desktop/AISIEVAL.xlsx")
D = "2026-08-17"
shutil.copy2(P, P.replace(".xlsx", f".backup-{time.strftime('%Y%m%d-%H%M%S')}.xlsx"))
wb = openpyxl.load_workbook(P)
ws = wb[wb.sheetnames[0]]
hdr = [c.value for c in ws[1]]
ix = {h: i + 1 for i, h in enumerate(hdr) if h}

MATRIX = {("C1", "Substantive"): "Proportionate", ("C1", "Partial"): "Under-response (gap)",
          ("C1", "Acknowledged"): "Under-response (gap)", ("C1", "None"): "Accountability gap (no action)",
          ("C2", "Substantive"): "Proportionate", ("C2", "Partial"): "Proportionate",
          ("C2", "Acknowledged"): "Under-response (gap)", ("C2", "None"): "Accountability gap (no action)"}

BAT56 = (f"Channel A battery run {D}. (1) company newsroom/blog; (2) the model's own launch card — "
         "GPT-5.6 System Card (2026-07-09), retrieved in full; (3) company safety hub "
         "(deploymentsafety.openai.com/gpt-5-6); (4) official company posts; (5) open web search to locate "
         "primary sources only. Admissible response located at (2): the launch card is the canonical response "
         "location for a pre-deployment evaluation per §8 source 2.")

BATM5 = (f"Channel A battery run {D}. (1) anthropic.com/news; (2) the model's own launch card — Claude Fable 5 "
         "& Claude Mythos 5 System Card, retrieved in full and searched for the 14% figure and for "
         "research-sabotage mitigation language; (3) Anthropic safeguards / RSP pages; (4) official company "
         "posts; (5) open web search to locate primary sources only. Battery exhausted. The card engages the "
         "finding at length but documents no safeguard, deployment change or mitigation attributable to it; the "
         "'testing is ongoing / we plan to iterate' language elsewhere in the card belongs to §3.3.1 cyber "
         "safeguards, not to research sabotage.")

BATFT = (f"Channel A battery run {D}. (1) openai.com newsroom/blog — no post referencing the paper or the "
         "pointwise-undetectable attack class; (2) no launch card applies (the finding targets an API surface, "
         "not a model release); (3) developers.openai.com fine-tuning docs and safety-best-practices — the "
         "moderation_checks event type is a pre-existing feature, not a change attributable to this finding; "
         "(4) official company posts, incl. the 2025-09-12 US CAISI / UK AISI update, which covers biological-"
         "misuse safeguards and product configuration, not fine-tuning pointwise defences; (5) open web search "
         "to locate primary sources only. ONE CANDIDATE FOUND AND NOT ADMITTED PENDING A RULING: the "
         "self-serve fine-tuning wind-down (notified 2026-05-07, phased to 2027-01-06) removes the exact attack "
         "surface this finding exercised. It is already admitted as a Partial response on four FAR.AI rows, so "
         "this row is inconsistent with them either way. Held at None pending §8 ruling — see Notes.")

RULING = (f"[CROSS-ROW INCONSISTENCY — RAISED {D}] This row's coding cannot be settled without a rulebook "
          "ruling, because the corpus currently treats the same company action two ways. OpenAI's self-serve "
          "fine-tuning wind-down (notified 2026-05-07, phased to 2027-01-06, confirmed at "
          "developers.openai.com/api/docs/deprecations as a wind-down of the platform itself rather than of "
          "specific models) removes the attack surface this finding exercised. Four FAR.AI rows — "
          "FARAI-2024-08-JAI1, FARAI-2024-10-ALI1, FARAI-2024-10-JAI1, FARAI-2025-02a-JAI2 — already cite that "
          "wind-down as a Partial access-restriction response at lags of 459 to 641 days with No explicit "
          "attribution. This row and PALISADE-2024-12-JAI1 sit at None on the same surface. Only two outcomes "
          "are internally consistent: (a) the wind-down is admissible, and these two rows become Partial, "
          "moving them out of the accountability gap; or (b) it is not admissible — no safety rationale is "
          "stated and no evaluator is referenced, so the causal link is the coder's inference — and the four "
          "FAR.AI rows fall to None, enlarging the gap. Held at None, the conservative direction for the "
          "paper's claim, pending that decision. Rows correctly NOT affected: FARAI-2025-02a-JAI3 (Anthropic "
          "surface) and UKAISI-2025-08-JAI2 (open-weight models), where the OpenAI platform is not the surface.")

CH = {
    "JOINT-2026-07-CYB1": dict(
        AL="Substantive", BAT=BAT56,
        RESP=("OpenAI designated GPT-5.6 as High capability in Cybersecurity under its Preparedness Framework and "
              "activated safeguards for the launch, with UK AISI's cyber evaluation reported in the card among the "
              "external cyber-capability inputs. VERBATIM (GPT-5.6 system card §1): \"Under our Preparedness "
              "Framework, we are treating Sol, Terra and Luna as High capability in both Cybersecurity and "
              "Biological and Chemical risk. None of them reach our High threshold in AI Self-Improvement. We have "
              "implemented a tailored set of safeguards, adapted to each model's capability profile, to "
              "sufficiently minimize the associated risks.\" The card also records an access restriction: at the US "
              "government's request the launch began as \"a limited preview for a small group of trusted partners "
              "whose participation has been shared with the government, before releasing more broadly.\""),
        NOTE=("None -> Substantive; Proportionality Accountability gap -> Proportionate (C1 + Substantive). The "
              "original log reasoned that the card \"reports the capability result but no company action tied "
              "specifically to the capability level itself\". That is inconsistent with how every other "
              "capability finding in the corpus is treated, including OPENAI-2026-06-SELF-CYB1 — Irregular's "
              "FrontierCyber result on the same model, in the same card, coded Substantive on exactly this "
              "designation. Under the adjudication rule recorded across the High-capability bloc, a Preparedness "
              "designation is Substantive where the card cites this finding among the designation's inputs and "
              "safeguards activated as a result. Both hold: UK AISI's cyber-range result is reported under "
              "\"External Evaluations for Cyber Capabilities – UK AISI\" and the card states safeguards were "
              "implemented to sufficiently minimize the associated risks. Treating two external cyber-capability "
              "evaluations of the same model in the same document differently was the defect.")),
    "UKAISI-2026-06-ALI6": dict(
        AL="Acknowledged", BAT=BATM5,
        RESP=("Anthropic reproduced UK AISI's research-sabotage result in the Claude Fable 5 & Mythos 5 system "
              "card and set out why it matters, but documented no safeguard or deployment change attributable to "
              "it. VERBATIM: \"[Claude Mythos 5] actively continues compromising safety research 14% of the time, "
              "compared with 2% of the time for Mythos Preview and 1% of the time for both Opus 4.7 and Opus "
              "4.8.\" The card adds \"Though we are uncertain about the implications of these findings for "
              "deployment risk, this evaluation is still a potentially useful signal. Therefore, the increased "
              "rate of continuing safety research compromise behaviors from [Claude Mythos 5] is noteworthy\", "
              "then lists three risk mechanisms and the evaluation's substantial limitations."),
        NOTE=("None -> Acknowledged; Proportionality Accountability gap -> Under-response (gap) (C2 + "
              "Acknowledged). Outside the C1 headline population, so the headline is unaffected. None was wrong "
              "on the corpus's own convention for company-published findings: where the company's card reports "
              "the evaluator's result and engages the underlying problem but specifies no action, the corpus "
              "codes Acknowledged, not None — see OPENAI-2026-03-SELF-ALI2 and JOINT-2026-02-ALI2. This card does "
              "more than report the number: it argues why the result is noteworthy and enumerates three risk "
              "mechanisms. That is recognition of the underlying problem. Confirmed no action attaches: the "
              "card's forward-looking \"testing is ongoing... we plan to iterate on this system over time\" sits "
              "in §3.3.1 on cyber safeguards, a different subject.")),
    "UKAISI-2025-02-JAI1": dict(BAT=BATFT, EXTRA_NOTE=RULING),
}

log = []
for r in range(2, ws.max_row + 1):
    fid = ws.cell(r, 1).value
    if fid not in CH:
        continue
    c = CH[fid]
    old_al = ws.cell(r, ix["Action Level"]).value
    old_pr = ws.cell(r, ix["Proportionality"]).value
    if c.get("AL"):
        ws.cell(r, ix["Action Level"]).value = c["AL"]
        sev = ws.cell(r, ix["Severity (C1/C2) majority"]).value
        ws.cell(r, ix["Proportionality"]).value = MATRIX[(sev, c["AL"])]
    if c.get("RESP"):
        ws.cell(r, ix["Company Response"]).value = c["RESP"]
        ws.cell(r, ix["Response Date"]).value = ws.cell(r, ix["Publication Date"]).value
        ws.cell(r, ix["Lag (days)"]).value = 0
        ws.cell(r, ix["Attribution"]).value = "Explicit attribution"
        ws.cell(r, ix["Channel A Evidence"]).value = ws.cell(r, ix["Source URL"]).value
    ws.cell(r, ix["Sources Checked (channel A)"]).value = c["BAT"]
    nt = ws.cell(r, ix["Notes"]).value or ""
    add = (f"\n\n[ACTION LEVEL RE-VERIFIED {D}] " + c["NOTE"]) if c.get("NOTE") else ""
    if c.get("EXTRA_NOTE"):
        add += "\n\n" + c["EXTRA_NOTE"]
    ws.cell(r, ix["Notes"]).value = (nt + add).strip()
    log.append((fid, old_al or "None", ws.cell(r, ix["Action Level"]).value,
                old_pr, ws.cell(r, ix["Proportionality"]).value))
wb.save(P)
for f, a, b, c2, d2 in log:
    print(f"  {f:<24} {a:<13} -> {b:<13} | {c2:<32} -> {d2}")

ws2 = openpyxl.load_workbook(P, data_only=True)[wb.sheetnames[0]]
A = [r for r in range(2, ws2.max_row + 1) if str(ws2.cell(r, ix["Action Trackable?"]).value or "") == "yes"]
H = [r for r in A if str(ws2.cell(r, ix["Severity (C1/C2) majority"]).value or "") == "C1"]
c = collections.Counter(str(ws2.cell(r, ix["Proportionality"]).value or "") for r in H)
gap, und = c["Accountability gap (no action)"], c["Under-response (gap)"]
print(f"\nC1 headline n={len(H)}")
print(f"  gap          {gap}/{len(H)} = {gap/len(H):.1%}   (was 92/147 = 62.6%)")
print(f"  falls short  {gap+und}/{len(H)} = {(gap+und)/len(H):.1%}  (was 116/147 = 78.9%)")
print("  action level:", dict(collections.Counter(str(ws2.cell(r, ix["Action Level"]).value or "") for r in A).most_common()))
bad = [ws2.cell(r, 1).value for r in A
       if MATRIX.get((str(ws2.cell(r, ix["Severity (C1/C2) majority"]).value or ""),
                      str(ws2.cell(r, ix["Action Level"]).value or "")))
       not in (None, str(ws2.cell(r, ix["Proportionality"]).value or ""))]
print("  matrix mismatches:", len(bad), bad[:5])
