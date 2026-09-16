#!/usr/bin/env python3
"""Build the 10-minute deck as a Keynote-editable .pptx.

Everything is a real PowerPoint object -- text frames, tables, pictures -- so Keynote can edit
type, move boxes and restyle. No text is baked into an image except inside the chart PNGs
themselves. Speaker notes go in the notes slide, which Keynote imports as presenter notes.
"""
import os, collections, datetime
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

CH = os.path.expanduser("~/Desktop/AISIEVAL_charts")
OUT = os.path.expanduser("~/Desktop/AISIEVAL_10min.pptx")
W, H = Inches(13.333), Inches(7.5)                      # 16:9

INK = RGBColor(0x14, 0x18, 0x1D)
INK2 = RGBColor(0x3A, 0x44, 0x4F)
MUTED = RGBColor(0x66, 0x70, 0x7C)
ACCENT = RGBColor(0x1F, 0x6F, 0xB2)
RED = RGBColor(0xC6, 0x2F, 0x1F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GROTESK = "Helvetica Neue"

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def tb(slide, x, y, w, h, text, size=18, color=INK, bold=False, font=None,
       align=PP_ALIGN.LEFT, spacing=1.0, anchor=MSO_ANCHOR.TOP, italic=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        f = r.font
        f.size = Pt(size)
        f.color.rgb = color
        f.bold = bold
        f.italic = italic
        f.name = font or GROTESK
    return box


def pic(slide, name, x, y, maxw, maxh):
    """Place a chart, scaled to fit the box and centred in it."""
    path = os.path.join(CH, name)
    iw, ih = Image.open(path).size
    s = min(maxw / iw, maxh / ih)
    w, h = int(iw * s), int(ih * s)
    return slide.shapes.add_picture(path, int(x + (maxw - w) / 2), int(y + (maxh - h) / 2),
                                    width=w, height=h)


def rule(slide, x, y, w, color=RGBColor(0xDC, 0xE2, 0xE9), h=Emu(9525)):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def slide_frame(no, clock, dur, title, land):
    """Standard content slide: number/clock rail, title, landing line."""
    s = prs.slides.add_slide(BLANK)
    tb(s, Inches(0.55), Inches(0.30), Inches(1.6), Inches(0.5),
       f"{no}", size=15, color=ACCENT, bold=True)
    tb(s, Inches(0.9), Inches(0.33), Inches(3.0), Inches(0.4),
       f"{clock}  ·  {dur}", size=11.5, color=MUTED, font="Menlo")
    tb(s, Inches(0.55), Inches(0.72), Inches(12.2), Inches(0.95), title,
       size=33, color=INK, bold=True, spacing=0.95)
    tb(s, Inches(0.55), Inches(1.72), Inches(12.2), Inches(0.75), land,
       size=17, color=INK2, spacing=1.15)
    rule(s, Inches(0.55), Inches(2.52), Inches(12.23))
    return s


# ----------------------------------------------------------------- 1
s = prs.slides.add_slide(BLANK)
tb(s, Inches(0.7), Inches(0.55), Inches(11.9), Inches(0.5),
   "EVALUATING THE EVALUATORS", size=13, color=ACCENT, bold=True, font="Menlo")
tb(s, Inches(0.7), Inches(1.05), Inches(11.9), Inches(1.5),
   "80% of significant-risk findings\ndraw no proportionate response",
   size=46, color=INK, bold=True, spacing=0.95)
tb(s, Inches(0.7), Inches(2.62), Inches(11.9), Inches(0.6),
   "Of 147 findings that name a company and were graded the more serious of two risk levels, "
   "118 drew no response proportionate to what they found.", size=17, color=INK2, spacing=1.15)
pic(s, "00_title_hero.png", Inches(0.7), Inches(3.35), Inches(11.93), Inches(3.35))
tb(s, Inches(0.7), Inches(6.85), Inches(11.9), Inches(0.4),
   "1,013 findings · 441 reports · 47 evaluators · Jan 2023 – Jul 2026",
   size=11.5, color=MUTED, font="Menlo")
notes(s, "0:00–0:50 — Open on the claim and let the funnel sit there. Do NOT explain tiers yet: "
         "the audience needs to want the method before you give it to them.\n\n"
         "The number: 118 of 147. 91 drew nothing at all, 27 drew something too weak.")

# ----------------------------------------------------------------- 2
s = slide_frame(2, "0:50", "1:10", "What earns a finding a place in the denominator",
                "1,013 findings → 193 where both the evaluation and a company action are traceable → "
                "147 graded significant risk by a three-model severity ensemble. Two filters, both "
                "pre-registered, neither applied by hand.")
pic(s, "20_accountability_pipeline_funnel.png", Inches(0.55), Inches(2.75), Inches(7.6), Inches(4.3))
pic(s, "07_tier_distribution.png", Inches(8.35), Inches(2.75), Inches(4.4), Inches(2.1))
pic(s, "17_severity_classification.png", Inches(8.35), Inches(4.95), Inches(4.4), Inches(2.1))
notes(s, "0:50–2:00 — The credibility slide; the denominator is the whole argument.\n\n"
         "Say out loud what is EXCLUDED: reassuring nulls, benchmark scores with no dangerous "
         "threshold crossed, comparative rankings, anonymised models. That kills the "
         "'you picked the bad cases' objection.\n\n"
         "Severity is a 3-model ensemble majority, not our judgement — that kills "
         "'you graded these yourself'.")

# ----------------------------------------------------------------- 3
s = slide_frame(3, "2:00", "1:15", "Scrutiny is scaling. That is not the bottleneck.",
                "47 institutions, 441 reports — 310 findings from government AI safety institutes, "
                "703 from third-party evaluators — and the volume rises steeply through 2025–26.")
pic(s, "13_institution_type_tree.png", Inches(0.55), Inches(2.75), Inches(7.5), Inches(4.3))
pic(s, "23_corpus_growth_by_tier.png", Inches(8.25), Inches(2.75), Inches(4.5), Inches(2.1))
pic(s, "12_evaluator_scope.png", Inches(8.25), Inches(4.95), Inches(4.5), Inches(2.1))
notes(s, "2:00–3:15 — Tree runs left-to-right so institution names stay readable.\n\n"
         "KEY POINT: the headline holds on the government-only subset too — 41/54 = 75.9% fall "
         "short — so it is not an artifact of pooling AISIs with independent labs.\n\n"
         "Then land the setup: scrutiny is growing, and the next slide shows response is not "
         "growing with it.")

# ----------------------------------------------------------------- 4
s = slide_frame(4, "3:15", "1:30", "Most significant-risk findings go unanswered",
                "91 no action · 27 under-response · 29 proportionate. 80.3% fall short of a response "
                "proportionate to the severity — 91 drew nothing at all, 27 drew something too weak.")
pic(s, "04_headline_outcome_distribution.png", Inches(0.55), Inches(2.75), Inches(6.0), Inches(4.3))
pic(s, "10_proportionality_by_severity.png", Inches(6.75), Inches(2.75), Inches(6.0), Inches(4.3))
notes(s, "3:15–4:45 — Your longest slide. Take the full 90 seconds; do not rush it.\n\n"
         "Define 'proportionate' out loud from the severity x response matrix, and note it is "
         "SEVERITY-RELATIVE: for a significant-risk finding only a substantive response qualifies, "
         "whereas for the lower severity a partial one does.\n\n"
         "Say plainly: this measures the CONTENT of the public response. It does not establish that "
         "the action was implemented, or that it reduced the risk.")

# ----------------------------------------------------------------- 5
s = slide_frame(5, "4:45", "1:10", "When companies do respond, what does it look like?",
                "35 substantive, 16 partial, 18 acknowledged. 57 responses cite the evaluator "
                "explicitly, 12 act without saying why. Median lag: 0 days.")
pic(s, "08_action_level_distribution.png", Inches(0.55), Inches(2.75), Inches(4.0), Inches(4.3))
pic(s, "18_attribution_distribution.png", Inches(4.68), Inches(2.75), Inches(4.0), Inches(4.3))
pic(s, "14_response_lag_distribution.png", Inches(8.8), Inches(2.75), Inches(4.0), Inches(4.3))
notes(s, "4:45–5:55 — The zero-day median is a FINDING, not a data error.\n\n"
         "Responses cluster in the launch system card, because that is where pre-deployment "
         "evaluations get answered. A response arriving after a finding is published is the rare case.\n\n"
         "The 12 'acted without saying why' rows are why Attribution is a separate column from "
         "Action Level — the dataset stops conflating silence with inaction.")

# ----------------------------------------------------------------- 6
s = slide_frame(6, "5:55", "0:45", "Almost nothing reaches binding policy",
                "173 no policy uptake · 19 non-binding uptake · ONE binding policy action across the "
                "entire corpus.")
pic(s, "09_policy_level_distribution.png", Inches(2.4), Inches(2.75), Inches(8.5), Inches(4.3))
notes(s, "5:55–6:40 — CHECK BEFORE PRESENTING: it is ONE binding action, not two.\n\n"
         "It is the June 2026 Commerce export-control order taking Fable 5 and Mythos 5 offline "
         "globally. The second candidate was re-sourced and downgraded: voluntary compliance at "
         "ONCD/OSTP request is not an enforceable instrument.\n\n"
         "This is the weakest link in the whole pipeline. Findings reach companies sometimes; "
         "they reach binding policy essentially never.")

# ----------------------------------------------------------------- 7
s = slide_frame(7, "6:40", "1:30", "Unevenly distributed, and narrowing slowly",
                "Shortfall varies sharply by developer and by evaluator access. Over time it declines: "
                "100% → 90% → 82% → 74% — still roughly three in four.")
pic(s, "15b_shortfall_rate_by_year.png", Inches(0.55), Inches(2.75), Inches(6.3), Inches(4.3))
pic(s, "19_gap_rate_by_developer.png", Inches(7.05), Inches(2.75), Inches(5.7), Inches(2.1))
pic(s, "16_gap_rate_by_access_type.png", Inches(7.05), Inches(4.95), Inches(5.7), Inches(2.1))
notes(s, "6:40–8:10 — On the headline metric the trend is DOWNWARD. Do not claim it is getting "
         "worse. The honest line: the shortfall is narrowing but remains the norm.\n\n"
         "Caveats first: 2023 is n=5, 2024 n=20, and 2026 is partial (cutoff 31 July). Fig 19 covers "
         "developers with n>=5 and the smallest bars are fragile.\n\n"
         "Some of the decline is composition, not behaviour change — later years hold more "
         "pre-deployment evaluations, which are answered in the launch card by construction.\n\n"
         "NOTE: the left chart is on the headline metric (falls short). The two right-hand charts plot "
         "the no-action rate, which is the lower number — say which you are quoting.")

# ----------------------------------------------------------------- 8
s = slide_frame(8, "8:10", "1:50", "Why you should believe it, and what would change it",
                "All 69 rows with a documented response were re-verified against the source document; "
                "9 were recoded, in both directions. The weak link is not detection — it is that "
                "nobody is obliged to answer.")
pic(s, "00_title_hero.png", Inches(1.3), Inches(2.85), Inches(10.7), Inches(3.0))
tb(s, Inches(0.55), Inches(6.15), Inches(12.2), Inches(1.0),
   "A response register would be cheap, and would make this measurable by default.",
   size=22, color=RED, bold=True)
notes(s, "8:10–10:00 — Corrections went BOTH ways (under-coded and over-coded), which is the "
         "evidence the review is not tuned to produce a bigger number.\n\n"
         "Every negative coding is backed by a completed, dated five-source search battery — an "
         "absence is only claimed where the search was exhausted.\n\n"
         "State the live limitation: the two corpus halves differ in Channel A search depth, so "
         "80.3% is an upper bound.\n\n"
         "Then the ask, once: a response register. Close on the hero you opened with, so the number "
         "lands twice. Then stop talking.")

# ----------------------------------------------------------------- 9 (reserve)
s = prs.slides.add_slide(BLANK)
tb(s, Inches(0.55), Inches(0.45), Inches(12.2), Inches(0.6),
   "Held in reserve — for Q&A", size=28, color=INK, bold=True)
tb(s, Inches(0.55), Inches(1.12), Inches(12.2), Inches(0.4),
   "Nine figures kept off the running order, each mapped to the question it answers.",
   size=15, color=MUTED)
rows = [("\"Don't they answer before launch?\"", "05_pre_vs_post_deployment_response_rate"),
        ("\"Is the gap worse in some risk domains?\"", "22_domain_x_outcome_heatmap"),
        ("\"Do they at least answer the serious ones?\"", "21_severity_x_action_heatmap"),
        ("\"Is this just one company?\"", "02_findings_per_model_developer"),
        ("\"Do big evaluators get better treatment?\"", "24_evaluator_volume_vs_gap"),
        ("\"What do evaluators actually test?\"", "11_domain_distribution"),
        ("\"Who publishes the most?\"", "01_findings_per_institution"),
        ("\"Do evaluators get real access?\"", "03_findings_per_access_type"),
        ("Simpler alternative to slide 3's growth chart", "06_findings_per_year")]
tbl = s.shapes.add_table(len(rows) + 1, 2, Inches(0.55), Inches(1.75),
                         Inches(12.23), Inches(5.0)).table
tbl.columns[0].width = Inches(6.4)
tbl.columns[1].width = Inches(5.83)
for j, head in enumerate(("If asked…", "Show this figure")):
    c = tbl.cell(0, j)
    c.text = head
    p = c.text_frame.paragraphs[0]
    p.runs[0].font.size = Pt(13)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = WHITE
    p.runs[0].font.name = GROTESK
    c.fill.solid()
    c.fill.fore_color.rgb = ACCENT
for i, (q, f) in enumerate(rows, start=1):
    for j, val in enumerate((q, f + ".png")):
        c = tbl.cell(i, j)
        c.text = val
        p = c.text_frame.paragraphs[0]
        p.runs[0].font.size = Pt(12.5)
        p.runs[0].font.color.rgb = INK2
        p.runs[0].font.name = "Menlo" if j else GROTESK
notes(s, "Not presented. Keep this slide in the file so the figures are one click away in Q&A.\n\n"
         "05_pre_vs_post is the one you are most likely to need.")

prs.save(OUT)
print(f"wrote {OUT}")
print(f"  {len(prs.slides.__iter__.__self__._sldIdLst)} slides · 16:9 · "
      f"{os.path.getsize(OUT):,} bytes")
