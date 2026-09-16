#!/usr/bin/env python3
"""How is the OpenAI fine-tuning attack surface treated across the corpus?

If the 2026-05-07 self-serve fine-tuning wind-down is admitted as an access-restriction
response, it must be applied consistently to every finding whose attack surface it removes
-- otherwise UKAISI-2025-02-JAI1 gets a coding the neighbouring rows do not.
"""
import openpyxl, os, re, datetime

P = os.path.expanduser("~/Desktop/AISIEVAL.xlsx")
ws = openpyxl.load_workbook(P, data_only=True)[openpyxl.load_workbook(P).sheetnames[0]]
hdr = [c.value for c in ws[1]]
ix = {h: i + 1 for i, h in enumerate(hdr) if h}


def norm(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


R = [{h: norm(ws.cell(r, ix[h]).value) for h in ix} for r in range(2, ws.max_row + 1)
     if ws.cell(r, 1).value]
tier = lambda r: ("A" if r["Action Trackable?"] == "yes"
                  else "B" if r["Eval? (trackable)"] == "yes" else "C")

FT = re.compile(r'fine.?tun', re.I)
hits = [r for r in R if FT.search(r["Finding"] + " " + r["Models / Systems"] + " " + r["Report Title"])]
oa = [r for r in hits if re.search(r'openai|gpt|o[134]\b', r["Models / Systems"] + r["Finding"], re.I)]
print(f"rows mentioning fine-tuning: {len(hits)} · of which OpenAI-surface: {len(oa)}\n")
print(f"{'Finding ID':<28} {'T':<2} {'sev':<4} {'Action Level':<13} {'Attribution':<24} cites wind-down?")
print("-" * 118)
for r in sorted(oa, key=lambda r: r["Publication Date"]):
    wd = "YES" if re.search(r'deprecat|wind.?down|2027-01-06|2026-05-07',
                            r["Company Response"] + r["Channel A Evidence"] + r["Notes"], re.I) else "-"
    print(f"{r['Finding ID']:<28} {tier(r):<2} {r['Severity (C1/C2) majority']:<4} "
          f"{r['Action Level']:<13} {r['Attribution']:<24} {wd}")
print("\n--- the four FAR.AI rows in full ---")
for r in oa:
    if r["Finding ID"].startswith("FARAI"):
        print(f"\n{r['Finding ID']}  AL={r['Action Level']}  attr={r['Attribution']}  lag={r['Lag (days)']}")
        print(f"   finding : {r['Finding'][:230]}")
        print(f"   response: {r['Company Response'][:340]}")
