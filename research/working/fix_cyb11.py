#!/usr/bin/env python3
"""UKAISI-2026-04-CYB11: Attribution 'No explicit attribution' -> 'Explicit attribution'.

Col 18 Explicit attribution = the responding company directly referenced the finding, report,
evaluation result or evaluating institution. OpenAI's own GPT-5.5 cybersecurity page does
exactly that, and the row's own response cell already quotes it.
"""
import openpyxl, os, shutil, time, collections, datetime

P = os.path.expanduser("~/Desktop/AISIEVAL.xlsx")
D = "2026-08-17"
shutil.copy2(P, P.replace(".xlsx", f".backup-{time.strftime('%Y%m%d-%H%M%S')}.xlsx"))
wb = openpyxl.load_workbook(P)
ws = wb[wb.sheetnames[0]]
hdr = [c.value for c in ws[1]]
ix = {h: i + 1 for i, h in enumerate(hdr) if h}

NOTE = (f"[ATTRIBUTION CORRECTED {D}] \"No explicit attribution\" -> \"Explicit attribution\". The row's own "
        "response cell quotes OpenAI's GPT-5.5 cybersecurity page naming the evaluating institution: \"We "
        "collaborated with the UK AI Security Institute (UK AISI) on pre-deployment testing for cyber "
        "capabilities.\" Col 18 Explicit attribution covers a company directly referencing the finding, report, "
        "evaluation result OR evaluating institution, so the previous value was contradicted by the row's own "
        "evidence. Found while auditing which positive codings actually rest on an undocumented causal inference: "
        "at lag -7 with the evaluator named in the company's own document, this row is a properly attributed "
        "pre-deployment response and does not belong in that group at all.")

for r in range(2, ws.max_row + 1):
    if ws.cell(r, 1).value != "UKAISI-2026-04-CYB11":
        continue
    old = ws.cell(r, ix["Attribution"]).value
    ws.cell(r, ix["Attribution"]).value = "Explicit attribution"
    nt = ws.cell(r, ix["Notes"]).value or ""
    ws.cell(r, ix["Notes"]).value = (nt + "\n\n" + NOTE).strip()
    print(f"  UKAISI-2026-04-CYB11  {old!r} -> 'Explicit attribution'")
wb.save(P)


def norm(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


ws2 = openpyxl.load_workbook(P, data_only=True)[wb.sheetnames[0]]
R = [{h: norm(ws2.cell(r, ix[h]).value) for h in ix} for r in range(2, ws2.max_row + 1)
     if ws2.cell(r, 1).value]
tier = lambda r: ("A" if r["Action Trackable?"] == "yes"
                  else "B" if r["Eval? (trackable)"] == "yes" else "C")
A = [r for r in R if tier(r) == "A"]
H = [r for r in A if r["Severity (C1/C2) majority"] == "C1"]
print("\nAttribution (Tier A):", dict(collections.Counter(r["Attribution"] for r in A).most_common()))

bloc = [r for r in H if r["Attribution"] == "No explicit attribution"
        and r["Action Level"] in ("Substantive", "Partial", "Acknowledged")]
c = collections.Counter(r["Proportionality"] for r in H)
gap, und, n = c["Accountability gap (no action)"], c["Under-response (gap)"], len(H)
CAND_IN_H = ["METR-2025-06-ALI1", "METR-2025-06-ALI2", "METR-2025-06-ALI3",
             "UKAISI-2025-02-JAI1", "PALISADE-2024-12-JAI1"]
print(f"\ngenuine inference-only bloc: {len(bloc)} C1 rows")
for r in sorted(bloc, key=lambda r: -int(r["Lag (days)"] or 0)):
    print(f"   {r['Finding ID']:<26} lag {r['Lag (days)']:>4}  {r['Action Level']:<12} {r['Proportionality']}")
print(f"\nCURRENT           gap {gap}/{n} = {gap/n:.1%} · falls short {gap+und}/{n} = {(gap+und)/n:.1%}")
gA = gap - len(CAND_IN_H)
print(f"OPTION A (uphold) gap {gA}/{n} = {gA/n:.1%} · falls short {gap+und}/{n} = {(gap+und)/n:.1%}")
gB = gap + len(bloc)
uB = und - sum(1 for r in bloc if r["Proportionality"] == "Under-response (gap)")
print(f"OPTION B (reject) gap {gB}/{n} = {gB/n:.1%} · falls short {gB+uB}/{n} = {(gB+uB)/n:.1%}")
