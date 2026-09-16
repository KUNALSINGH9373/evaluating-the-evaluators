#!/usr/bin/env python3
"""Fix the two Tier A rows miscoded Access Type = Aggregate.

Both name exactly one model, so 'Aggregate' cannot apply -- the value was inherited from the
report's cross-model trend framing and not updated when the finding was split per named model
(the split is what made these rows Tier A in the first place). Three siblings from the same
report on the same models are Post-deployment, including CYB2 on the identical model as CYB4.
"""
import openpyxl, os, shutil, time, collections, datetime

P = os.path.expanduser("~/Desktop/AISIEVAL.xlsx")
D = "2026-08-17"
shutil.copy2(P, P.replace(".xlsx", f".backup-{time.strftime('%Y%m%d-%H%M%S')}.xlsx"))
wb = openpyxl.load_workbook(P)
ws = wb[wb.sheetnames[0]]
hdr = [c.value for c in ws[1]]
ix = {h: i + 1 for i, h in enumerate(hdr) if h}

NOTE = (f"[ACCESS TYPE CORRECTED {D}] Access Type \"Aggregate\" -> \"Post-deployment\". Aggregate cannot apply to "
        "this row: it names exactly one model. The value was inherited from the report's cross-model trend framing "
        "(\"How Fast Is Autonomous AI Cyber Capability Advancing?\") and was not updated when the report's findings "
        "were split per named model under §5 -- and that split is precisely what makes this row Tier A, since the "
        "named-model test is what Aggregate rows fail. Corrected to match its siblings from the identical report on "
        "the identical models: UKAISI-2026-05-CYB1 (Tier B), UKAISI-2026-05-CYB2 (Tier A, same model as CYB4) and "
        "UKAISI-2026-05-GOV1 (Tier C) are all Post-deployment. Both models were publicly released before the "
        "2026-05-13 publication date, so no pre-deployment access is implied. This also empties the Aggregate "
        "bucket in Tier A, which was distorting figure 16 (gap rate by access type) with a 2-row cell.")

TARGET = {"UKAISI-2026-05-CYB3", "UKAISI-2026-05-CYB4"}
done = []
for r in range(2, ws.max_row + 1):
    fid = ws.cell(r, 1).value
    if fid not in TARGET:
        continue
    old = ws.cell(r, ix["Access Type"]).value
    ws.cell(r, ix["Access Type"]).value = "Post-deployment"
    nt = ws.cell(r, ix["Notes"]).value or ""
    ws.cell(r, ix["Notes"]).value = (nt + "\n\n" + NOTE).strip()
    done.append((fid, old))
wb.save(P)
for f, o in done:
    print(f"  {f:<24} {o!r} -> 'Post-deployment'")

ws2 = openpyxl.load_workbook(P, data_only=True)[wb.sheetnames[0]]
tier = lambda r: ("A" if str(ws2.cell(r, ix["Action Trackable?"]).value or "") == "yes"
                  else "B" if str(ws2.cell(r, ix["Eval? (trackable)"]).value or "") == "yes" else "C")
ct = collections.Counter((str(ws2.cell(r, ix["Access Type"]).value or ""), tier(r))
                         for r in range(2, ws2.max_row + 1) if ws2.cell(r, 1).value)
print("\nAccess Type x Tier after the fix:")
for at in ("Pre-deployment", "Post-deployment", "Mixed", "Aggregate", "N/A"):
    print(f"  {at:<18} A={ct[(at,'A')]:>4}  B={ct[(at,'B')]:>4}  C={ct[(at,'C')]:>4}")
