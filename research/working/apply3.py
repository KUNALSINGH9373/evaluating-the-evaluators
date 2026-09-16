import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
P='/Users/kunalsingh/Desktop/v11_FINAL.xlsx'
wb=openpyxl.load_workbook(P); ws=wb['v11 findings']
del wb['Legend']
lg=wb.create_sheet('Legend')
B=Font(bold=True)
rows=[
 ('FINAL v11 CANDIDATE SET — post-review, post §1 evaluation-vs-benchmark pass (2026-08-15)',None,None,None),
 (None,)*4,
 ('Colour','Status','Rows','What it means / what to do'),
 ('Yellow','REVIEW-TIER-A',89,'ACCEPTED and Tier A — these drive the accountability analysis. Verify the verbatim quote and the split/club decision, then mark Sign-off.'),
 ('Orange','REVIEW-DATE',49,'Publication date missing, year-only, or inferred rather than read. §6.2 requires a source-verified date shared by all rows of a Report ID. Note: 34 of these also carry placeholder Report/Finding IDs that must be regenerated once dated.'),
 ('Green','READY',430,'Passed every check. Ready to stage as-is.'),
 (None,None,568,'ACTIVE TOTAL'),
 (None,)*4,
 ('Red','Removed sheet',35,'Disqualified — moved off the main sheet, not deleted. 29 previously EXCLUDED + 6 newly disqualified by the §1 evaluation-vs-benchmark test.'),
 (None,)*4,
 ('Tier split (active rows)',None,None,None),
 (None,'Tier A',94,None),(None,'Tier B',293,None),(None,'Tier C',181,None),
 (None,)*4,
 ('§1 evaluation-vs-benchmark test — applied 2026-08-15',None,None,None),
 (None,'The test','Who RAN the evaluation and asserts the result — not who published the document, and not whose questions were used.',None),
 (None,'KEEP','The document states that a third party evaluated the model, ran the tests, or raised the concern. The developer is then only the host.',None),
 (None,'DROP','The document states that the developer used a third party\'s benchmark, dataset, scaffold or framework to evaluate its own model. That is the developer\'s evaluation, not a joint exercise.',None),
 (None,'Also DROP','No named evaluator — "additional external partners", "an external group". Anonymous attribution fails the named-evaluator test.',None),
 (None,'Note','Benchmark co-authorship is not contribution. Gray Swan co-developed ART with UK AISI; UK AISI is therefore NOT an author of findings produced by running ART.',None),
 (None,)*4,
 ('What changed on 2026-08-15',None,None,None),
 (None,'Reinstated',1,'Row for OPENAI-2025-04-SELF-CYB1 (Pattern Labs / o3). Prior exclusion was wrong — the conclusion is in the card, later in §3.9.3 than the extractor read. Retiered A->B.'),
 (None,'Newly disqualified',6,'2 Meta-run Agentic Misalignment rows misattributed to Apollo; 1 Anthropic-organised IPI competition; 2 Mythos rows with unnamed evaluators; 1 Anthropic self-review of Irregular incidents.'),
 (None,'Roster hold cleared',10,'Irregular (8) and Signature Science (2). Verified in source that both RAN the evaluations. Tagged developer-commissioned for a sensitivity analysis.'),
 (None,'Attribution corrected',5,'SHADE-Arena x4: Scale AI added as named co-author. Sonnet 4.6 Gray Swan x1: UK AISI removed (benchmark co-development only), scope government-AISI -> third-party-evaluator.'),
 (None,'Content corrected',2,'Muse Spark / Gray Swan: Meta\'s AgentDojo 11.7% figure removed from a Gray Swan row. GPT-4o / METR: quote swapped to METR\'s own §4.1 assessment.'),
 (None,'Retained by ruling',2,'SLEIGHT-Bench (4:1 Anthropic:Redwood authorship). Contrast SHADE-Arena, a stated three-way collaboration — the stronger case on the same rule.'),
 (None,)*4,
 ('Not yet done',None,None,None),
 (None,'Severity (C1/C2) ensemble — needs API keys; all severity columns are blank',None,None),
 (None,'Channel A/B/C response batteries — required for Tier A rows only',None,None),
 (None,'Proportionality — derived, needs severity + Action Level',None,None),
 (None,'Roster policy: whether developer-commissioned vendors get their own Scope value for the headline statistic (10 rows tagged)',None,None),
]
for r in rows: lg.append(list(r))
for cell in ('A1','A3','B3','C3','D3','A11','A16','A23','A31'):
    lg[cell].font=B
lg['A3'].fill=lg['B3'].fill=lg['C3'].fill=lg['D3'].fill=PatternFill('solid',start_color='D9D9D9',end_color='D9D9D9')
for c,col in zip(('FFF2CC','FCE4D6','E2EFDA'),(4,5,6)):
    lg.cell(col,1).fill=PatternFill('solid',start_color=c,end_color=c)
lg.cell(9,1).fill=PatternFill('solid',start_color='F4CCCC',end_color='F4CCCC')
for col,w in (('A',34),('B',24),('C',10),('D',96)): lg.column_dimensions[col].width=w
for r in range(1,lg.max_row+1): lg.cell(r,4).alignment=Alignment(wrap_text=True,vertical='top')
wb.move_sheet('Legend', offset=len(wb.sheetnames)-1-wb.sheetnames.index('Legend'))
wb.save(P)
print(wb.sheetnames, ws.max_row-1)
