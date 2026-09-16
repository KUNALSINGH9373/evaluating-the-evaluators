import openpyxl, collections
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
P='/Users/kunalsingh/Desktop/v11_FINAL.xlsx'
wb=openpyxl.load_workbook(P); ws=wb['v11 findings']
hdr=[c.value for c in ws[1]]; ix={h:i+1 for i,h in enumerate(hdr) if h}
S=lambda r,k,v: ws.cell(r,ix[k]).__setattr__('value',v)
G=lambda r,k: ws.cell(r,ix[k]).value

# ---- 5. clear the roster hold: verified that these evaluators RAN the evaluations ----
EV={'Irregular':"GPT-5.6 Preview System Card §9.1.2.5: 'Irregular evaluated GPT-5.6 Sol across three offensive cybersecurity evaluation suites ... Irregular found that GPT-5.6 Sol has on-par or slightly stronger offensive-cyber capabilities than GPT-5.5.'",
    'Signature Science':"Deep Research System Card: 'Biosecurity experts from Signature Science probed the Pre-Mitigation deep research model for conceptual ideation of designing novel biological threats.'"}
for r in range(2,ws.max_row+1):
    if G(r,'Status')!='REVIEW-ROSTER': continue
    inst=str(G(r,'Institution') or '')
    ev=EV['Signature Science'] if 'Signature' in inst else EV['Irregular']
    tierA = G(r,'Action Trackable?')=='yes'
    S(r,'Status','REVIEW-TIER-A' if tierA else 'READY')
    S(r,'Review note',("Roster hold cleared 2026-08-15. The §1 test is who ran the evaluation and asserts the result, not how the organisation reached the roster. Verified in source: "+ev+
       (" Tier A — verify the verbatim quote and the split/club decision, then sign off." if tierA else " Tier B/C, no further action.")))
    S(r,'Notes',(G(r,'Notes') or '').split(' Extracted')[0].rstrip('. ')+
       ". Roster hold cleared 2026-08-15 under the §1 evaluation-vs-benchmark test — evaluator ran the evaluation and asserts the result. Developer-commissioned vendor with no independent publishing venue; flag retained in Tags for a sensitivity analysis. TODO severity ensemble + Channel A/B/C.")
    S(r,'Tags',(G(r,'Tags') or '')+';developer-commissioned')

# ---- 6. move EXCLUDED + REMOVED to a Removed sheet, then delete from main ----
kill=[r for r in range(2,ws.max_row+1) if G(r,'Status') in ('EXCLUDED','REMOVED')]
if 'Removed' in wb.sheetnames: del wb['Removed']
rem=wb.create_sheet('Removed')
rem.append(hdr)
for r in kill:
    rem.append([ws.cell(r,c).value for c in range(1,ws.max_column+1)])
for r in sorted(kill, reverse=True):
    ws.delete_rows(r)
print('removed', len(kill), '-> active', ws.max_row-1)

# ---- 7. restyle ----
FILL={'REVIEW-TIER-A':'FFF2CC','REVIEW-DATE':'FCE4D6','READY':'E2EFDA'}
for r in range(2,ws.max_row+1):
    f=PatternFill('solid',start_color=FILL[G(r,'Status')],end_color=FILL[G(r,'Status')])
    for c in range(1,ws.max_column+1): ws.cell(r,c).fill=f
ws.auto_filter.ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
ws.freeze_panes='D2'

# Removed-sheet styling
hf=PatternFill('solid',start_color='D9D9D9',end_color='D9D9D9')
for c in range(1,rem.max_column+1):
    rem.cell(1,c).fill=hf; rem.cell(1,c).font=Font(bold=True)
rf=PatternFill('solid',start_color='F4CCCC',end_color='F4CCCC')
for r in range(2,rem.max_row+1):
    for c in range(1,rem.max_column+1): rem.cell(r,c).fill=rf
rem.freeze_panes='D2'
for col,w in (('A',16),('B',70),('C',12),('D',30),('E',26),('F',30)):
    rem.column_dimensions[col].width=w
rem.auto_filter.ref=f"A1:{get_column_letter(rem.max_column)}{rem.max_row}"

wb.save(P)
c=collections.Counter(G(r,'Status') for r in range(2,ws.max_row+1))
print(dict(c))
tier=collections.Counter(('A' if G(r,'Action Trackable?')=='yes' else 'B' if G(r,'Eval? (trackable)')=='yes' else 'C') for r in range(2,ws.max_row+1))
print(dict(tier))
