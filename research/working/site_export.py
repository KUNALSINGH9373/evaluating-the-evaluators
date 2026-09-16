#!/usr/bin/env python3
"""Export AISIEVAL.xlsx to v11.csv in the repo, and report the numbers the site hardcodes.

build_data.py consumes a CSV with the 40-column schema. The workbook is the source of truth, so
the CSV is a derived artifact -- regenerate it, never hand-edit it.
"""
import openpyxl, os, csv, collections, datetime, statistics

SRC = os.path.expanduser("~/Desktop/AISIEVAL.xlsx")
REPO = os.path.expanduser("~/evaluating-the-evaluators")
OUT = os.path.join(REPO, "v11.csv")


def norm(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


ws = openpyxl.load_workbook(SRC, data_only=True)["AISIEVAL"]
hdr = [norm(c.value) for c in ws[1]]
rows = []
for r in range(2, ws.max_row + 1):
    rec = {h: norm(ws.cell(r, i).value) for i, h in enumerate(hdr, 1) if h}
    if rec.get("Finding ID"):
        rows.append(rec)

with open(OUT, "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=[h for h in hdr if h])
    w.writeheader()
    w.writerows(rows)
print(f"wrote {OUT}  ({len(rows)} findings x {len([h for h in hdr if h])} columns, "
      f"{os.path.getsize(OUT)/1e6:.2f} MB)")

# ---- the numbers app.js / index.html currently hardcode -------------------------------
tier = lambda r: ("A" if r["Action Trackable?"] == "yes"
                  else "B" if r["Eval? (trackable)"] == "yes" else "C")
A = [r for r in rows if tier(r) == "A"]
H = [r for r in A if r["Severity (C1/C2) majority"] == "C1"]
sev = collections.Counter(r["Severity (C1/C2) majority"] for r in rows)
VOTES = ("Sonnet5 vote", "GPT-5.5 vote", "Gemini3.1 vote")
unan = split = novote = 0
for r in rows:
    v = [r.get(c, "") for c in VOTES]
    if not all(v):
        novote += 1
        continue
    if len(set(v)) == 1:
        unan += 1
    else:
        split += 1
prop = collections.Counter(r["Proportionality"] for r in H)
n = len(H)
short = prop["Accountability gap (no action)"] + prop["Under-response (gap)"]
print(f"""
--- values to substitute into the site ---
  findings              {len(rows)}          (was 456)
  reports               {len({r['Report ID'] for r in rows if r['Report ID']})}          (was 211)
  severity C1           {sev['C1']}          (was 155)
  severity C2           {sev['C2']}          (was 301)
  unanimous 3-0         {unan}  = {unan/len(rows)*100:.0f}%   (was 415 = 91%)
  split 2-1             {split}  = {split/len(rows)*100:.0f}%   (was 41 = 9%)
  rows with no 3 votes  {novote}
  accountability set    {len(A)}
  headline population   {n}
  falls short           {short}/{n} = {short/n*100:.1f}%
  no response at all    {prop['Accountability gap (no action)']}/{n} = {prop['Accountability gap (no action)']/n*100:.1f}%
  date range            {min(r['Publication Date'] for r in rows if r['Publication Date'])} to {max(r['Publication Date'] for r in rows if r['Publication Date'])}
""")
