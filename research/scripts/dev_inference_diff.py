#!/usr/bin/env python3
"""Where do the two developer inferences disagree, and which rows are ambiguous?

Figure 19 makes per-company claims. Both the chart generator and the independent verifier infer
the developer by pattern-matching free text, and they disagree, so neither can be treated as
authoritative. This lists every disagreement and every row matching more than one vendor.
"""
import openpyxl
import dataset_source, datetime, re, collections, os

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
ws = dataset_source.sheet()
hdr = [c.value for c in ws[1]]
ix = {h: i + 1 for i, h in enumerate(hdr) if h}


def n(v):
    if v is None: return ""
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime("%Y-%m-%d")
    return str(v).strip()


R = [{h: n(ws.cell(r, ix[h]).value) for h in ix} for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
tier = lambda r: ("A" if r["Eval? (trackable)"] == "yes" and r["Action Trackable?"] == "yes"
                  else "B" if r["Eval? (trackable)"] == "yes" else "C")
H = [r for r in R if tier(r) == "A" and r["Severity (C1/C2) majority"] == "C1"]

# --- as implemented in charts.py -------------------------------------------------------
DEVMAP = [("OpenAI", ("gpt", "o1-", "o3", "o4-", "chatgpt", "codex", "operator", "gpt-oss")),
          ("Anthropic", ("claude", "opus", "sonnet", "haiku", "mythos", "fable")),
          ("Google", ("gemini", "gemma", "deepmind")),
          ("Meta", ("llama", "prompt-guard", "promptguard", "muse spark")),
          ("DeepSeek", ("deepseek",)), ("Alibaba (Qwen)", ("qwen",)), ("Mistral", ("mistral",)),
          ("Moonshot", ("kimi", "moonshot")), ("xAI", ("grok",)), ("Z.ai", ("glm",)),
          ("Thinking Machines", ("inkling",))]
def dev_chart(r):
    s = (r["Models / Systems"] + " " + r["Report Title"]).lower()
    for d, ks in DEVMAP:
        if any(k in s for k in ks): return d
    return "Other / unnamed"

# --- as implemented in verify_charts.py ------------------------------------------------
VMAP = [(r'gpt|chatgpt|codex|\bo[134]\b|operator|sora', "OpenAI"),
        (r'claude|opus|sonnet|haiku|fable|mythos', "Anthropic"),
        (r'gemini|gemma|palm', "Google"), (r'llama|prompt.?guard', "Meta"),
        (r'grok', "xAI"), (r'deepseek', "DeepSeek"), (r'qwen', "Alibaba (Qwen)"),
        (r'mistral|magistral', "Mistral"), (r'nemotron', "NVIDIA")]
def dev_verify(r):
    s = r["Models / Systems"] + " " + r["Report Title"]
    for pat, name in VMAP:
        if re.search(pat, s, re.I): return name
    return "Other / unnamed"

diff = [(r, dev_chart(r), dev_verify(r)) for r in H if dev_chart(r) != dev_verify(r)]
print(f"headline population {len(H)} · rows where the two inferences disagree: {len(diff)}\n")
for r, a, b in diff:
    print(f"  {r['Finding ID']:<30} charts.py={a:<16} verifier={b}")
    print(f"      models : {r['Models / Systems'][:100]}")
    print(f"      report : {r['Report Title'][:100]}")

# --- rows that legitimately match more than one vendor ---------------------------------
print("\nrows matching MORE THAN ONE vendor (the inference silently picks the first):")
multi = 0
for r in H:
    s = r["Models / Systems"] + " " + r["Report Title"]
    hit = sorted({name for pat, name in VMAP if re.search(pat, s, re.I)})
    if len(hit) > 1:
        multi += 1
        print(f"  {r['Finding ID']:<30} {hit}")
        print(f"      {r['Models / Systems'][:100]}")
print(f"  -> {multi} multi-vendor rows out of {len(H)}")

for label, fn in (("charts.py", dev_chart), ("verifier", dev_verify)):
    c = collections.Counter(fn(r) for r in H)
    c.pop("Other / unnamed", None)
    print(f"\n{label:<10} n>=5: " + " · ".join(f"{d} {v}" for d, v in c.most_common() if v >= 5)
          + f"   (unattributed {sum(1 for r in H if fn(r)=='Other / unnamed')})")
