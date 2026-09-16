#!/usr/bin/env python3
"""Append the 168 merge-ready sweep findings to the V13 sheet, in place."""
import sys, os, json, collections
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import dataset_source as ds
import openpyxl

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
M = json.load(open(os.path.join(ROOT, "logs/window_extension/merge/merge_ready.json")))

wb = openpyxl.load_workbook(ds.WORKBOOK)
ws = wb[ds.SHEET]
hdr = [ds.norm(c.value) for c in ws[1]]
ix = {h: i + 1 for i, h in enumerate(hdr) if h}
start = ws.max_row
while start > 1 and not ws.cell(start, 1).value: start -= 1

before = start - 1
seen = {ds.norm(ws.cell(r, 1).value) for r in range(2, start + 1)}
row = start + 1
written = 0
for m in M:
    assert m["Finding ID"] not in seen, f"ID collision {m['Finding ID']}"
    seen.add(m["Finding ID"])
    # tier is DERIVED; write the two source columns and let every reader derive from them.
    tier = "A" if m["eval"] == "yes" and m["action"] == "yes" else "B" if m["eval"] == "yes" else "C"
    assert tier == m["tier"], (m["Finding ID"], tier, m["tier"])
    v = {
        "Finding ID": m["Finding ID"], "Report ID": m["Report ID"],
        "Institution": m["Institution"], "Institution Type": m["Institution Type"],
        "Report Title": m["title"], "Publication Date": m["date"], "Domain": m["Domain"],
        "Models / Systems": m["models"], "Access Type": m["Access Type"], "Source URL": m["url"],
        "Finding": m["finding"], "Finding Quote": m["quote"],
        "Eval? (trackable)": m["eval"],
        # S9 col 36: blank on Tier C rows, else yes/no.
        "Action Trackable?": ("" if tier == "C" else m["action"]),
        "Finding Type": m["Finding Type"], "Scope": m["Scope"],
        "Sources Checked (channel A)": f"[ADDED {m['src']} · window-extension sweep 2026-08-29] "
                                       + (m.get("reasoning") or "")[:600],
    }
    for k, val in v.items():
        if k in ix: ws.cell(row, ix[k]).value = val
    row += 1; written += 1

wb.save(ds.WORKBOOK)
print(f"appended {written} rows   {before} -> {before + written}")
