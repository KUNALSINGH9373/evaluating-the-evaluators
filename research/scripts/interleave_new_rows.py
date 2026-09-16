#!/usr/bin/env python3
"""Insert the 167 sweep rows at their correct positions, shifting V12 rows down.

Two placement rules, in order:

  1. If the new row belongs to a report V12 ALREADY holds, it goes immediately after that report's
     existing rows. This is what makes the 41 split reports contiguous again — the report keeps one
     position in the sheet, determined by where V12 already put it.
  2. Otherwise the new row's whole report group is inserted at the first point where the V12
     sequence has reached an older publication date.

V12's RELATIVE order is preserved exactly: no V12 row ever moves past another V12 row. Only its
absolute row number changes, as new rows slot in above it. V12's own 3 date-order irregularities
are left as they are — they are the author's sheet, not mine to normalise.
"""
import sys, os, re, collections
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import dataset_source as ds
import openpyxl

def natkey(s):
    return [int(t) if t.isdigit() else t for t in re.split(r'(\d+)', s or "")]

V12P = "/Users/kunalsingh/evaluating-the-evaluators/AISI  Eval Findings.xlsx"
src = openpyxl.load_workbook(V12P, data_only=True, read_only=True)["AISIEVAL_V12"]
srows = list(src.iter_rows(values_only=True))
shdr = [str(c).strip() if c is not None else "" for c in srows[0]]
V12_IDS = [str(r[shdr.index("Finding ID")]).strip() for r in srows[1:] if r[shdr.index("Finding ID")]]
V12_SET = set(V12_IDS)

wb = openpyxl.load_workbook(ds.WORKBOOK)
ws = wb[ds.SHEET]
hdr = [ds.norm(c.value) for c in ws[1]]
ncol = len(hdr)
i_id, i_rep, i_date = hdr.index("Finding ID"), hdr.index("Report ID"), hdr.index("Publication Date")

body = [[ws.cell(r, c).value for c in range(1, ncol + 1)]
        for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value]
snap = {ds.norm(r[i_id]): tuple(ds.norm(v) for v in r) for r in body}
byid = {ds.norm(r[i_id]): r for r in body}

old = [byid[f] for f in V12_IDS]                          # V12, in V12's own sequence
new = [r for r in body if ds.norm(r[i_id]) not in V12_SET]
assert len(old) + len(new) == len(body), (len(old), len(new), len(body))

rep_of  = lambda r: ds.norm(r[i_rep])
date_of = lambda r: ds.norm(r[i_date])[:10]

# --- rule 1: new rows joining a report V12 already has -------------------------------------
v12_reps = {rep_of(r) for r in old}
attach = collections.defaultdict(list)
standalone = []
for r in new:
    (attach[rep_of(r)] if rep_of(r) in v12_reps else standalone).append(r)
for k in attach:
    attach[k].sort(key=lambda r: natkey(ds.norm(r[i_id])))

# --- rule 2: brand-new reports, grouped, ordered newest first ------------------------------
groups = collections.defaultdict(list)
for r in standalone:
    groups[rep_of(r)].append(r)
for k in groups:
    groups[k].sort(key=lambda r: natkey(ds.norm(r[i_id])))
pending = sorted(groups.items(), key=lambda kv: (date_of(kv[1][0]), natkey(kv[0])), reverse=True)

# --- weave ---------------------------------------------------------------------------------
out = []
gi = 0
for idx, r in enumerate(old):
    # insert any brand-new report whose date is newer than the V12 row we are about to write
    while gi < len(pending) and date_of(pending[gi][1][0]) > date_of(r):
        out.extend(pending[gi][1]); gi += 1
    out.append(r)
    # last row of this report in the V12 sequence? then append its new siblings here
    last = (idx + 1 == len(old)) or rep_of(old[idx + 1]) != rep_of(r)
    if last and rep_of(r) in attach:
        out.extend(attach.pop(rep_of(r)))
while gi < len(pending):
    out.extend(pending[gi][1]); gi += 1

assert not attach, f"unattached report groups: {list(attach)}"
assert len(out) == len(body), (len(out), len(body))
after = {ds.norm(r[i_id]): tuple(ds.norm(v) for v in r) for r in out}
assert after == snap, "cell contents changed — aborting"

seq = [ds.norm(r[i_id]) for r in out if ds.norm(r[i_id]) in V12_SET]
assert seq == V12_IDS, "V12 relative order was disturbed — aborting"

for i, row in enumerate(out):
    for c, v in enumerate(row, start=1):
        ws.cell(i + 2, c).value = v
wb.save(ds.WORKBOOK)
print(f"wrote {len(out)} rows")
print(f"V12 relative order preserved exactly : {seq == V12_IDS}")
print(f"new rows attached to existing reports: {sum(1 for r in new if rep_of(r) in v12_reps)}")
print(f"brand-new report groups inserted     : {len(pending)}")
